# -*- coding: utf-8 -*-
"""Minimal python-docx helpers for a conference paper (A4, Calibri, header-shaded tables, page numbers).

Built for the English edition of a Korean hwpx 발제문 (2026-09). Mirrors the hwpx builder's shapes:
title block -> right-aligned event lines -> "About this paper" box -> numbered sections with
bullet paragraphs and 2-3 column tables whose multi-paragraph cells get "- " markers.

Usage:
    from docx_paper_lib import Paper
    p = Paper()                      # A4, 2.3 cm margins, Calibri 11, centered "- N -" page footer
    p.P("Title", size=17, bold=True, color=p.NAVY, align="center")
    p.H1("1. Section"); p.B("bullet text"); p.T(["Category", "Content"], [["a", ["p1", "p2"]]], [1.0, 4.0])
    p.save(out_docx)
Convert with export_pdf.ps1 (Word COM, wdFormatPDF=17):  powershell -File export_pdf.ps1 -Src x.docx -Out x.pdf
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}


class Paper:
    NAVY = RGBColor(0x1F, 0x3A, 0x5F)
    HDR_FILL = "DCE6F1"

    def __init__(self, font="Calibri", east_asia="Malgun Gothic", body_pt=11, margin_cm=2.3):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(margin_cm)
        self.body_w = 21.0 - 2 * margin_cm
        st = self.doc.styles["Normal"]
        st.font.name = font
        st.font.size = Pt(body_pt)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.12
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.add_run("- ")
        fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"; r.append(t); fld.append(r)
        fp._p.append(fld)
        fp.add_run(" -")

    @staticmethod
    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def P(self, text, size=11, bold=False, color=None, align=None, after=4, before=0, italic=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        if align:
            p.alignment = _ALIGN[align]
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if color:
            r.font.color.rgb = color
        return p

    def H1(self, text):
        p = self.P(text, size=13.5, bold=True, color=self.NAVY, before=10, after=4)
        p.paragraph_format.keep_with_next = True
        return p

    def H2(self, text):
        p = self.P(text, size=12, bold=True, color=self.NAVY, before=6, after=3)
        p.paragraph_format.keep_with_next = True
        return p

    def B(self, text, marker="●"):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(marker + "  "); r.font.size = Pt(9); r.font.color.rgb = self.NAVY
        r = p.add_run(text); r.font.size = Pt(11)
        return p

    @staticmethod
    def _cell(cell, content, bold=False, center=False, size=10.5):
        paras = content if isinstance(content, list) else [content]
        if isinstance(content, list) and len(paras) >= 2:
            paras = [("- " + x if not x.startswith(("※", "- ")) else x) for x in paras]
        cell.paragraphs[0].text = ""
        for i, t in enumerate(paras):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            if center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(t); r.font.size = Pt(size); r.bold = bold

    def T(self, header, rows, widths, align=None, after=8):
        """header: list of column titles; rows: list of rows whose cells are str or [paragraphs];
        widths: relative column widths; align: 'C'/'L' per column (default first column centered)."""
        align = align or ("C" + "L" * (len(header) - 1))
        tbl = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        ws = [Cm(self.body_w * w / sum(widths)) for w in widths]
        for ci, h in enumerate(header):
            c = tbl.rows[0].cells[ci]
            self._cell(c, h, bold=True, center=True)
            self.shade(c, self.HDR_FILL)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                self._cell(tbl.rows[ri + 1].cells[ci], val, bold=(ci == 0), center=(align[ci] == "C"))
        for row in tbl.rows:
            for ci, c in enumerate(row.cells):
                c.width = ws[ci]
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(max(after - 4, 0))
        return tbl

    def box(self, title, body, size=10.5):
        """Single-column framed box: shaded title row + body row (used for 'About this paper')."""
        tbl = self.doc.add_table(rows=2, cols=1)
        tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._cell(tbl.rows[0].cells[0], title, bold=True, center=True); self.shade(tbl.rows[0].cells[0], self.HDR_FILL)
        self._cell(tbl.rows[1].cells[0], body, size=size)
        for row in tbl.rows:
            row.cells[0].width = Cm(self.body_w)
        self.P("", after=2)
        return tbl

    def save(self, path):
        self.doc.save(path)
        return path
