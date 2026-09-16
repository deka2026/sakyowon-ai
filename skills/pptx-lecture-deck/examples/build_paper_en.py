# -*- coding: utf-8 -*-
"""English edition of the conference paper (docx via python-docx), mirroring the final Korean hwpx/PDF."""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
HDR_FILL = "DCE6F1"

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin = sec.bottom_margin = Cm(2.3)
sec.left_margin = sec.right_margin = Cm(2.3)
BODY_W = 21.0 - 4.6

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.12

# footer page number
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("- ")
f1 = OxmlElement("w:fldSimple"); f1.set(qn("w:instr"), "PAGE")
rr = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = "1"; rr.append(tt); f1.append(rr)
fp._p.append(f1)
fp.add_run(" -")


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def P(text, size=11, bold=False, color=None, align=None, after=4, before=0, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if align: p.alignment = align
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    return p


def H1(text):
    p = P(text, size=13.5, bold=True, color=NAVY, before=10, after=4)
    p.paragraph_format.keep_with_next = True
    return p


def H2(text):
    p = P(text, size=12, bold=True, color=NAVY, before=6, after=3)
    p.paragraph_format.keep_with_next = True
    return p


def B(text, marker="●"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(marker + "  "); r.font.size = Pt(9); r.font.color.rgb = NAVY
    r = p.add_run(text); r.font.size = Pt(11)
    return p


def cell_write(cell, content, bold=False, center=False, size=10.5):
    paras = content if isinstance(content, list) else [content]
    if isinstance(content, list) and len(paras) >= 2:
        paras = [("- " + x if not x.startswith(("※", "- ")) else x) for x in paras]
    cell.paragraphs[0].text = ""
    for i, t in enumerate(paras):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(t); r.font.size = Pt(size); r.bold = bold


def T(header, rows, widths, align="CL", after=8):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    ws = [Cm(BODY_W * w / sum(widths)) for w in widths]
    for ci, h in enumerate(header):
        c = tbl.rows[0].cells[ci]
        cell_write(c, h, bold=True, center=True)
        shade(c, HDR_FILL)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.rows[ri + 1].cells[ci]
            cell_write(c, val, bold=(ci == 0), center=(align[ci] == "C"))
    for row in tbl.rows:
        for ci, c in enumerate(row.cells):
            c.width = ws[ci]
    # repeat header row
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(after - 4)
    return tbl


# ============================================================ Title block
P("Local Government-Led Citizen Renewable Energy: A Case from Korea", size=17, bold=True, color=NAVY,
  align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
P("— The Jeonnam-Gwangju Integrated Special City Transition Committee's Deliberations on "
  "\"Citizen-Sovereign Community Wealth Building\" —", size=11.5, bold=True, color=NAVY,
  align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
for t in ["Presentation at the 2026 Community Energy International Conference (Korea Energy Agency · REN21)",
          "Session: \"Community Energy Cases and Experiences from Around the World\"",
          "— Local Government-Led Community Energy and Local Climate Action in Korea —",
          "Thursday, 1 October 2026 | Korea Press Center, International Conference Hall, Seoul",
          "Presenter: Il-young Kim, Chairperson, Social Innovation Platform"]:
    P(t, size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT, after=1)
P("", after=4)

tbl = doc.add_table(rows=2, cols=1); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_write(tbl.rows[0].cells[0], "About this paper", bold=True, center=True); shade(tbl.rows[0].cells[0], HDR_FILL)
cell_write(tbl.rows[1].cells[0],
    "This paper introduces the deliberations of the Jeonnam-Gwangju Integrated Special City Transition Committee on the "
    "draft plan \"A Basic-Society Strategy Based on Renewable Energy Business — Citizen-Sovereign Community Wealth Building\" "
    "(July 2026), which is under review. It is a proposal at the discussion stage, not an adopted policy. The dedicated "
    "organizational structure and the budget and financing plan have not been finalized and are therefore not covered here. "
    "Instead, the paper concentrates on \"what, why, and by what principles\" — the program content and its design principles. "
    "It should be read not as a competitor to the national Sunlight Income Village program, but as a case of how a local "
    "government can extend and deepen that program.", size=10.5)
for row in tbl.rows: row.cells[0].width = Cm(BODY_W)
P("", after=2)

# ============================================================ 1
H1("1. Purpose and Scope of This Paper")
B("The subtitle of this conference is \"From Renewable Energy to Regional Prosperity.\" A renewable energy plant arriving in a "
  "region and its profits remaining in the region as residents' income and assets are two entirely different things. The "
  "Transition Committee's deliberations are precisely about how a local government can close that gap.")
B("The presentation takes the national Sunlight Income Village program as its point of departure, then introduces the content "
  "and principles of a proposal that widens the scope from administrative villages (ri) to townships (eup/myeon/dong) and "
  "urban cooperatives, and that goes beyond plant installation to design microgrids, village electrification and a circular "
  "economy.")
T(["Category", "Content"],
  [["Covered", ["Background and problem framing", "Design principles (asset-based social and solidarity economy) and reference models",
                "Objectives, and differences from the national Sunlight Income Village program", "The four programs and how they interlock",
                "Networks and scale; expected effects and remaining tasks"]],
   ["Not covered", ["Dedicated organizational structure (departments, teams, staffing, the inter-departmental working group, intermediary support organizations)",
                    "Budget and financing plan, unit subsidy levels, scale of public investment and grants",
                    "※ These were discussed by the Transition Committee but remain undecided, and are excluded for that reason"]]],
  [1.0, 4.0])

# ============================================================ 2
H1("2. Background — Why Does a Local Government Take On Renewable Energy?")
B("As electricity demand surges in the AI era, expanding renewable generation infrastructure has become unavoidable. The "
  "question is who receives the gains. The starting point of the Transition Committee's deliberations is to design the program "
  "so that the gains from renewable expansion flow back to citizens as assets and income.")
T(["Background", "Content"],
  [["A new era of power demand", "Expand renewable generation infrastructure to meet surging AI-era electricity demand, while channeling the gains into citizens' assets and income"],
   ["Jeonnam-Gwangju's conditions", ["A core region of three national mega-projects and an urban–rural region with large renewable potential",
                                      "Use the region's geographic, social and economic advantages so that the Integrated Special City consolidates a leading position in the AI economy, building the renewable ecosystem first and then spreading to urban and rural everyday economies"]],
   ["Problem framing — energy colony, or energy sovereign?", ["In centralized, corporate renewable systems the region bears the production burden while the profits flow out",
                                                                "Distributed (\"local production, local consumption\") renewable infrastructure eases the region's generation burden and gives it a fair share of the benefits as sovereign"]],
   ["Inheriting the assets of the two former regions", ["(Gwangju) about 1,300 social economy enterprises and the experience of citizens' solar cooperatives",
                                                         "(Jeonnam) renewable hubs such as offshore wind, village communities, the fishing-village revitalization program, and resident benefit-sharing schemes such as the Sunlight and Wind Pension",
                                                         "Combined into an \"urban–rural mutual prosperity, citizen-sovereign economy\""]]],
  [1.3, 3.7])

# ============================================================ 3
H1("3. Why Build Community Wealth with Renewable Energy? — From the Transition Committee's Q&A")
B("Of the questions and answers raised during the Committee's review, four that reveal the program's underlying logic are reproduced here.")
T(["Question", "Summary of answer"],
  [["Why build community wealth around renewable energy?",
    ["Electricity will be an essential good in the AI economy, and renewable generation is the most stable business model for producing it — making it the most desirable field for building and expanding community wealth. Investment and lending risk for the Citizen Fund and the Citizen Enterprise Fund can be regarded as close to nil.",
     "It is a stable model through which the community and local government can jointly build a governance-type social and solidarity economy (the Québec model), and it scales readily into housing, local industry and other fields as the energy transition proceeds."]],
   ["How does it differ from the central government's Sunlight Income Village program?",
    ["The Sunlight Income Village program is a first-stage program focused on installing power plants. Jeonnam-Gwangju pursues a pilot model aimed at microgrids for local production and consumption and the electrification of housing and the village economy, and builds its own social finance system to form and maintain village assets.",
     "The national program (centered on administrative villages) is extended to township and urban models while securing the same grid conditions (priority connection) and financial support as the national program, and the city works to have this extended model institutionalized in subsequent national calls."]],
   ["How does it differ from existing social and solidarity economy policy?",
    ["Existing support policy is moving from start-up support toward growth and stabilization support, but it does not build a value-chain ecosystem for a specific industry. What is different here is a model in which federations of village and township Community Interest Companies (solar cooperatives) and social and solidarity economy organizations specializing in component assembly, construction, operation and maintenance form an ecosystem with a firm mutual-trading market.",
     "Rural village cooperatives also need living-service businesses (care, culture, health) beyond the power plant, but are too small to sustain them. A service supply system pooled across several villages is therefore underwritten by the renewable energy business."]],
   ["Is a separate citizen-sovereignty AI needed alongside the city's AI platform?",
    ["AI trained on general knowledge tends to give plausible wrong answers in specialized fields and carries value-related risks. Citizen sovereignty, the social and solidarity economy and residents' autonomy rest heavily on tacit, experiential knowledge, so market-economy knowledge centered on quantitative data alone yields poor results from AI transformation.",
     "An independent, citizen-sovereignty domain AI trained on the values and tacit knowledge of the social and solidarity economy is therefore built to serve citizens and practitioners in the field."]]],
  [1.35, 3.65])

# ============================================================ 4
H1("4. Design Principles — Asset-Based Social and Solidarity Economy")
B("The theoretical background of this proposal is the Asset-Based Social and Solidarity Economy model. It combines three strands of "
  "theory: start from the assets a region already has rather than its deficits, and let the community own and circulate the wealth that used to leak out.")
T(["Theory", "Key content"],
  [["Asset-Based Community Development (ABCD)", "Endogenous development that starts from the assets a region holds (people, organizations, nature, space, finance) rather than from its deficits (Kretzmann & McKnight, 1993)"],
   ["Community Wealth Building", "Wealth that used to leak out of the region is retained through an \"anchor institution – local business – resident\" circulation, using five levers: ownership, employment, procurement, finance and land"],
   ["Social and Solidarity Economy (SSE) and the commons", "The reciprocal, solidarity-based economy of cooperatives and social enterprises (recognized by the UN, ILO and OECD), and the commons principle under which land, energy and data are democratically owned and managed by the community"]],
  [1.5, 3.5])
B("The theory is made concrete through five principles and seven instruments, and the renewable energy business is the first arena "
  "in which these instruments actually operate: idle public land (land), the village Community Interest Company (enterprise), the "
  "Citizen Fund (fund) and credit guarantees (credit) all combine within a single project.")
T(["Category", "Content"],
  [["Five principles", "Democratized asset ownership · reciprocity and solidarity · local circulation · democratic governance · self-reliance and resilience"],
   ["Seven instruments", "Community land trust · cooperative complex · Community Interest Company (CIC) · participatory budgeting · community fund · local currency · mutual credit"],
   ["Jeonnam-Gwangju's integrated strategy", ["An integrated strategy that develops these into \"renewable energy benefit sharing → Village Salary → community wealth.\" Beyond that, it aims to link renewable energy with rural basic income and the social and solidarity economy as a foundation for a citizens' \"basic society\"",
                                               "The principles are translated into program requirements — community bylaws, profit-sharing rules, trust agreements, management disclosure — that become the eligibility and selection criteria of each funding call"]]],
  [1.3, 3.7])

H2("□ Reference models at home and abroad")
B("The role of this proposal is to integrate proven models into a single ecosystem fitted to Jeonnam-Gwangju's urban–rural conditions. "
  "Guyang-ri in Yeoju, this conference's site-visit village, is the direct reference for the \"microgrid-linked pilot\" in the Jeonnam-Gwangju proposal.", marker="○")
T(["Category", "Model", "Reference point"],
  [["International", ["(Spain) Mondragón cooperative complex", "(USA) Cleveland Evergreen", "(UK) Preston Model", "(UK) Community shares"],
                      ["Inter-cooperative solidarity and internal finance", "Anchor-institution procurement and worker-owned firms", "Local procurement to stop wealth leaking out", "Non-transferable withdrawable capital, capped interest, one member one vote"]],
   ["Domestic", ["Shinan Sunlight and Wind Pension", "Citizens' Solar Power Cooperative (Gwangju); Wanju Local Food", "Ansan Federation of Solar Power Cooperatives"],
                 ["A precedent for resident benefit sharing", "Citizen-invested generation and a local food circulation system", "The federation model that binds village cooperatives together"]],
   ["Guyang-ri, Yeoju", "A microgrid-linked Sunlight Income Village", "Staged upgrade from the current 1 MW to 5 MW (the microgrid level rises with each 1 MW expansion); a target of KRW 500,000 per resident per month at 5 MW; a village education center"]],
  [0.8, 1.9, 2.3], align="CLL")

# ============================================================ 5
H1("5. Objectives")
B("The objective is to build an income system that flows back to residents and citizens, and the social and solidarity economy "
  "ecosystem beneath it. The quantitative targets are draft figures at the Transition Committee's review stage and have not been adopted.")
T(["Category", "Content"],
  [["Objective ① A citizen-sovereign income system", "Establish a locally tailored, citizen-sovereign income system that flows back to residents and citizens, designing participation routes not only for residents but also for non-resident stakeholders and urban citizens"],
   ["Objective ② \"Village Salary\" income activities", "Systematically support residents' \"Village Salary\" income activities based on community assets in renewables, tourism, agri-food processing and other fields, beginning with solar and wind income and diversifying into landscape trees, village tourism and processing"],
   ["Objective ③ Ecosystem infrastructure", "Build a Jeonnam-Gwangju \"asset-based social and solidarity economy\" ecosystem infrastructure grounded in the five principles and seven instruments"],
   ["An urban–rural mutual prosperity economy", "Redesign by combining Gwangju's social economy base with Jeonnam's renewable hubs and resident benefit sharing, widening the scope from villages (ri) → townships (eup/myeon/dong) → urban cooperatives as an extended version of the national program"],
   ["Quantitative targets (draft, cumulative 2026–2030)", ["250 township community organizations · 2,550 village community organizations supported",
                                                            "200 leading Community Interest Companies · 200 village economy infrastructure sites · 50 citizen enterprises",
                                                            "2,100 domain specialists trained · 400 active practitioners (cumulative)"]]],
  [1.4, 3.6])

# ============================================================ 6
H1("6. How Is It Different from the National Sunlight Income Village Program?")
B("The national program is a first stage focused on installing power plants. The Jeonnam-Gwangju proposal widens the scope on the "
  "same grid and finance conditions and aims at a pilot model that designs microgrids, village electrification and a circular economy from the outset.")
T(["Item", "National Sunlight Income Village", "Jeonnam-Gwangju model (draft)"],
  [["Scope", "Administrative villages (ri)", "Villages + township Residents' Autonomy Councils + cooperatives with urban members"],
   ["Business", "Village income limited to a solar power plant", "Same + support for locally tailored income models such as landscape trees, village tourism and agri-food processing"],
   ["Own contribution and finance", ["Construction loan secured by revenue rights with partial interest support", "Resident compensation payments and the Local Depopulation Response Fund count as own contribution"],
                                     ["Low-interest loans for the village's own share from a separate Citizen Fund", "Priority equity contributions from non-resident stakeholders count as own contribution"]],
   ["Grid and infrastructure", "Priority grid connection; ESS support where connection is not possible", "Same + microgrid infrastructure where grid connection is not possible"],
   ["Construction and O&M", "KEA-certified RESCOs (renewable energy service companies)", ["Citizen-invested, city-certified citizen enterprises participate as RESCOs", "Citizen enterprises in component manufacturing and village electrification are fostered, with a mutual-trading system"]],
   ["Network", "Selected villages must join the national network", ["Cooperatives (CICs) must join the management disclosure system", "Optional federation membership gives access to management support and to business services for living-service operations (legal, HR, accounting, marketing, joint purchasing)"]],
   ["Direction", "Limited to renewable generation (stage 1)", ["Energy transition villages through microgrids", "A growth- and stabilization-stage support system that systematizes the village and regional circular economy (at scale) from the start, beyond start-up support"]],
   ["Local government role", ["Government budget limited to interest subsidy and training", "Joint village–municipality applications; own contribution and grid connection burden officials and act as entry barriers"],
                              ["One-stop support at metropolitan and municipal level and a metropolitan governance control tower ease the burden on villages and officials", "Drive institutionalization of the township and urban extended model in national calls"]]],
  [1.0, 1.9, 2.1], align="CLL")

# ============================================================ 7
H1("7. Program Content — Four Interlocking Programs")
B("The program has four axes that interlock and circulate. Resident communities discover assets (①), Community Interest Companies "
  "and citizen enterprises create income (②), domain AI supports them (③), and finance recycles and reinvests the returns (④). "
  "The returns and funds from ④ are reinvested in new village communities under ① and new projects under ②. The period is July 2026 "
  "to June 2030 (draft), and the lead body is the Jeonnam-Gwangju Integrated Special City.")
T(["Program", "Key content", "Link to the next axis"],
  [["① Village Salary foundations", ["Residents' Autonomy Councils and village associations discover local assets and develop income models", "Village asset map · use rights to idle public land · legal entity · seed capital · democratic operating capacity"], "Communities with process records become eligible to apply to ② and ④"],
   ["② CIC and citizen enterprise ecosystem", ["One-stop start-up and growth support for CICs", "Village economy infrastructure (microgrids, processing, tourism)", "Fostering value-chain hub citizen enterprises"], "Enterprise activity data accumulates on the ③ platform; returns flow to ④"],
   ["③ AI-based Village Salary ecosystem", ["Domain-specific AI platform automating CIC and citizen enterprise operations", "Training Solidarity Intelligence AI Activists with a mandatory data-contribution structure"], "Track records and management disclosure build the trust base for the ④ funds"],
   ["④ Returns and community wealth", ["Returns flow to resident dividends · Citizen Fund reserves · reinvestment · welfare", "Credit guarantee + Citizen Fund Foundation + Citizen Enterprise Fund", "Charitable trusts lock community assets in for the long term"], "Funds are reinvested in new communities (①) and new projects (②)"]],
  [1.3, 2.4, 1.3], align="CLL")

H2("① Village Salary foundations — residents discover their own assets")
B("This is not a program that \"hands out\" a power plant; it supports the process by which resident communities survey and organize "
  "their own assets and derive an income model. Eligible applicants are resident communities such as Residents' Autonomy Councils and "
  "village associations, with bonus points for consortia that include non-resident stakeholders and citizens. Township community "
  "organizations, not only administrative villages, are eligible.", marker="○")
T(["Sub-program", "Content"],
  [["①-1 Local asset discovery and income model development", ["Activity + consulting package: ① local asset survey (a \"village asset map\") ② securing use rights to idle public land ③ support for establishing a legal entity ④ securing seed capital (resident equity, fund matching, compensation payments, Hometown Love Donations) ⑤ support for resident activities (operating costs, facilitation)",
                                                                  "Method: (stage 1) asset survey and resident organizing → (stage 2) use rights, entity and capital → (stage 3) income model. Follow-on calls are linked to stage results, with a consultant pool provided"]],
   ["①-2 Specialist consulting for income diversification", ["Fields: renewables (Sunlight Income Village type, agrivoltaic federation type) / village tourism (experience and stay infrastructure) / landscape trees (open-field smart farming) / agri-food processing (small smart processing plants)",
                                                              "Expert consultants matched by field → business plan upgrade → support for applying to central-government and metropolitan ecosystem calls → a priming track into ②"]],
   ["①-3 Democratic capacity building and process records", ["Training and workshops: community bylaws, profit-sharing rules, village vision, democratic decision-making (general meeting, board); designing the trust agreement between the village association and the CIC",
                                                              "Institutionalized process records: outputs, agreements and minutes are documented and made a mandatory eligibility requirement for ② and ④ ecosystem calls",
                                                              "A village consultant pool registered, trained and dispatched by field: legal, finance, energy, tourism, agriculture, governance"]]],
  [1.4, 3.6])

H2("② CIC and citizen enterprise ecosystem — the enterprises that create the Village Salary")
B("A two-layer enterprise ecosystem in which village Community Interest Companies (CICs) create income and citizen enterprises at "
  "metropolitan, regional and municipal level link the value chain. A Community Interest Company is a social and solidarity economy "
  "enterprise established to solve local problems and serve the public interest; a village solar cooperative is the typical example.", marker="○")
T(["Sub-program", "Content"],
  [["②-1 One-stop priming support for CICs", "One-stop support matching early-stage technology, management, finance and youth staffing → incubation → growth → the leading-enterprise track. A pool of specialist consultants and a consultation AI platform"],
   ["②-2 Growth support for leading CICs", "Selection criteria: ① Village Salary generated and planned ② local asset base ③ social and solidarity economy identity (democratic operation, recirculation of returns) ④ sustainability ⑤ process records on file (from ①-3)"],
   ["②-3 Village economy infrastructure", ["Composition: microgrids (per 100 households) + income infrastructure for agri-food processing, open-field smart landscape-tree farming and village tourism",
                                            "Requirements: ① a CIC established and operating in the village (mandatory) ② a trust agreement with the village association locking community assets in for the long term (mandatory) ③ operating and profit-sharing bylaws and a sustainability plan (mandatory) ④ ① activity base and process records",
                                            "Pilots from 2027: begin with some ten microgrid-linked pilot sites combining villages selected for the national Sunlight Income Village program with township and urban-cooperative models, then scale up — with technical support from the Korea Energy Agency and governance with KEPCO; other Village Salary models are also piloted",
                                            "Follow-up: trust-based community ownership; clawback where results fall short or funds are misused"]],
   ["②-4 Fostering citizen enterprises at metropolitan, regional and municipal level", ["Purpose: value-chain hub citizen enterprises — federation-type enterprises of several solar cooperatives, RESCO-type enterprises that install and operate facilities on villages' behalf and settle and distribute monthly returns, energy-saving appliance retailers, etc.",
                                                                                          "Requirements: ① evidence of value-chain business effectiveness ② social and solidarity economy character (democratic governance, profit reinvestment, local circulation) ③ an annual social contribution report ④ management disclosure"]]],
  [1.4, 3.6])

H2("□ Local production, local consumption — the logic of microgrids and village electrification")
B("Where the national Sunlight Income Village program is a first stage focused on plant installation, Jeonnam-Gwangju aims at microgrids "
  "for local production and consumption and at the electrification of housing and the village economy. Microgrid projects start from the "
  "beginning, a village-level electrification strategy is applied to housing and industry, and power-plant villages grow into \"energy transition villages\" over the long term.", marker="○")
T(["Category", "Content"],
  [["The order for handling grid issues", "① Secure priority grid connection on the same terms as the national program → ② where grid capacity is short, resolve with ESS support → ③ where connection is impossible, switch to microgrid infrastructure"],
   ["Staged upgrading (the Guyang-ri way)", "Raise the microgrid level with each capacity expansion, building the village's self-consumption, storage and management capacity together. Scale up from some ten pilot sites to village economy infrastructure (target: 100 microgrids)"],
   ["The regional meaning of distributed infrastructure", "In a core region of the three mega-projects, distributed (\"local production, local consumption\") rather than centralized, corporate infrastructure eases the region's generation burden. It is an alternative rationale for the AI-driven surge in energy demand, and a path for the region to benefit as sovereign rather than as an energy colony"],
   ["Remaining task", "Agreement on sharing the cost of microgrids (including ESS) among the grid operator (KEPCO), VPP operators, national government and local governments — building this governance comes first"]],
  [1.4, 3.6])

H2("③ AI-based Village Salary ecosystem — domain AI for citizen sovereignty")
B("Not general-purpose AI, but specialized AI trained on the values of the social and solidarity economy and the tacit knowledge of the "
  "field, supports village enterprises and practitioners. The key design feature is that practitioners are obliged to upload field data "
  "to the platform, creating a virtuous cycle in which data, AI and field activity grow together.", marker="○")
T(["Sub-program", "Content"],
  [["③-1 Building the Village Salary AI platform", ["Character: a domain-specific AI platform automating CIC and citizen enterprise operations (accounting, administration, marketing, disclosure) and providing information and participation guidance to workers and citizens",
                                                       "Six-step build: ① planning and data design ② field data accumulation and cleaning (practitioner labeling) ③ domain-specific model (LLM-based RAG and fine-tuning) ④ platform development ⑤ pilot with leading enterprises ⑥ full operation and upgrading (about 24 months)"]],
   ["③-2 Training and supporting Solidarity Intelligence AI Activists", ["Solidarity Intelligence AI Activists: field practitioners in Jeonnam-Gwangju's citizen-sovereign economy who contribute to accumulating the data that trains the specialized AI",
                                                                          "Training: completion of a domain-based AI course → activist certification",
                                                                          "Mandatory contribution: village consultants and field practitioners under ① and ② must upload process records, cases and data to the platform as a condition of activity and payment",
                                                                          "Virtuous cycle: practitioner data → platform AI upgrading → better enterprise automation → more field activity → more data"]]],
  [1.4, 3.6])

H2("④ Returns and community wealth — a social and solidarity finance system")
B("A social and solidarity finance system channels the returns from local assets into resident dividends, Citizen Fund reserves, "
  "reinvestment in new projects and welfare. How the fund formed from assets and returns is managed is the core task of this program, "
  "and the city builds a separate partnership with local financial institutions for it. Amounts and the mode of public investment are "
  "undecided, so only the structure and principles are introduced.", marker="○")
T(["Category", "Content"],
  [["[Public] Credit Guarantee Foundation", "Loan guarantees for CICs and citizen enterprises unlock lending from private and financial institutions, leveraging the guarantee reserve"],
   ["[Private ①] Citizen Fund Foundation (crowdfunding)", ["Donation type (prepaid pre-purchase) / loan type (low-interest lending) / investment type (lifetime equity — modeled on UK community shares: non-transferable, no capital gain, withdrawal at par only at the society's discretion, one member one vote)",
                                                            "The Citizen Fund is used only for the villages' own-contribution share of power plants"]],
   ["[Private ②] Citizen Enterprise Fund", "Equity and mezzanine investment in growth-stage citizen enterprises and CICs (scale-up capital), under impact-investment principles, with dividends recovered and recycled"],
   ["Governance and transparency", ["Citizen Fund: general meeting → board (citizen representatives + experts + public interest) → management office. The investment and loan committee applies a dual review of social value and financial soundness; information disclosure",
                                     "Citizen Enterprise Fund: an investor (LP) – manager (GP) structure with an investment review committee"]],
   ["The public role — from seed money to citizen capital", ["Public money is invested as seed capital at first; as citizen participation grows, the public stake is recovered in stages (repayment or transfer) and replaced by citizen capital, with recovered funds revolving",
                                                              "Grants are limited to non-recoverable items: interest-gap subsidies on loans, foundation and fund operations, the crowdfunding platform, education and publicity (following the Gyeonggi Province precedent of partial interest subsidy)",
                                                              "Roadmap: early years (1–2) investment-led → middle years (3–4) expanded participation and start of public recovery → long term, citizen-capital self-reliance"]],
   ["Long-term asset lock", "Support for establishing public-interest institutions (both publicly funded and private) that manage charitable trusts, citizen funding, the real transfer and long-term holding of public assets, and the reciprocal use of surplus assets"]],
  [1.5, 3.5])

# ============================================================ 8
H1("8. Beyond the Village — Networks, Federations and Living Services at Scale")
B("When villages bond into a solidarity ecosystem rather than each running its own power plant, income grows from a \"Village Salary\" into a regional economy.")
T(["Category", "Content"],
  [["National network and federations", ["Villages selected for the Sunlight Income Village program must join the national Sunlight Dividend Network, bonding them into a federation-level and nationwide ecosystem. The Ansan Federation of Solar Power Cooperatives is an example of the federation model",
                                          "The network could grow to about one million people on the basis of a target of 10,000 villages × 100 members per village"]],
   ["A mutual-trading market across the value chain", "Federations of village and township CICs (solar cooperatives) and social and solidarity economy organizations specializing in component assembly, construction, operation and maintenance form an ecosystem with a mutual-trading market. Unlike start-up-centered policy, this builds a value-chain ecosystem for a specific industry"],
   ["Living services at scale", "Rural village cooperatives need living-service businesses such as care, culture and health beyond the power plant, but their staff and markets are too small to sustain them. A supply system pooled across several villages is built, with the balance of supply and demand underwritten by the renewable energy business"],
   ["Transparency builds trust", "Cooperatives in selected villages must join the management disclosure system. Because track records and public management information are essential to trust in public finance and the Citizen Fund, systematic information management and disclosure are the rule from the start"]],
  [1.4, 3.6])

# ============================================================ 9
H1("9. Expected Effects and Remaining Tasks")
B("From building power plants to an economy that villages own and the region circulates — that is the direction of the Transition "
  "Committee's deliberations. Governance and institutionalization, and the final decisions on organization and budget, remain open.")
T(["Category", "Content"],
  [["Expected effects", ["Sustainable citizen income and community wealth by linking community economic returns to resident dividends, Citizen Fund reserves and reinvestment",
                         "Balanced development of Jeonnam-Gwangju through cooperation between residents and non-resident stakeholders",
                         "A circular regional economy less dependent on other regions and large corporations in manufacturing, distribution and finance",
                         "A tangible \"basic society\" foundation through returning part of the returns to the community"]],
   ["Remaining tasks", ["Governance: agreement on sharing microgrid costs among KEPCO, VPP operators, national government and local governments",
                        "Cooperation with central government: an MOU and task force with the Korea Energy Agency (the agency responsible for the Sunlight Income Village program), research on the Jeonnam-Gwangju pilot model, and institutionalization of the township and urban extended model in national calls",
                        "An integrated coordination system across overlapping departments (renewables, residents' autonomy, village autonomy, social and solidarity economy, microgrids)",
                        "The dedicated organization and the budget and financing method"]]],
  [1.2, 3.8])
P("< End >", align=WD_ALIGN_PARAGRAPH.RIGHT, before=6)

doc.save(OUT)
print("saved", OUT)
