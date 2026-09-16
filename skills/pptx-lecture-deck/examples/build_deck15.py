# -*- coding: utf-8 -*-
"""전남광주 시민주권 공동체 자산 형성 — 컨퍼런스 발표자료 빌더 (본문 15pt 이상판)"""
import sys, math
sys.path.insert(0, r"C:\Users\User\.claude\skills\pptx-lecture-deck\scripts")
from deck_lib import *
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

OUT = sys.argv[1]
SRC_NOTE = "자료: 전남광주통합특별시 인수위원회 검토용 사업계획(2026. 7.) · 조직체계·예산은 논의됐지만 미확정이라 제외"
AMBER = RGBColor(0xE0, 0x9F, 0x3E)
MIN = 15.0
CANDS = (17.0, 16.0, 15.0)


def header(s, num, title, lead):
    rect(s, 0, 0, SW, 1.12, fill=DARK)
    rect(s, MARGIN, 0.24, 0.66, 0.64, fill=LIME)
    tb = textbox(s, MARGIN, 0.30, 0.66, 0.56)
    para(tb.text_frame, f"{num:02d}", size=17, bold=True, color=DARK, first=True, align=PP_ALIGN.CENTER)
    tb = textbox(s, 1.42, 0.20, 11.4, 0.80)
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tb.text_frame, title, size=24, bold=True, color=WHITE, first=True, line=1.05)
    rect(s, 0, 1.12, SW, 0.78, fill=LIGHT)
    rect(s, 0, 1.12, 0.12, 0.78, fill=MID)
    tb = textbox(s, MARGIN, 1.16, 12.3, 0.70)
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tb.text_frame, lead, size=16, bold=True, color=DARK, first=True, line=1.12)


def footer(s):
    rect(s, MARGIN, 6.92, BODY_W, 0.03, fill=LINE)
    tb = textbox(s, MARGIN, 6.96, BODY_W, 0.45)
    para(tb.text_frame, SRC_NOTE, size=MIN, color=GRAY, first=True)


def card_h(heading, bullets, size, cw):
    tw = cw - 0.42
    hs = size + 2.5
    h = 0.16 + est_lines(heading, hs, cw - 0.4) * (hs * 1.20 / 72.0) + 0.15
    for b in bullets:
        t = b[1:] if b.startswith("*") else b
        h += est_lines("▸ " + t, size, tw) * (size * 1.22 / 72.0) + 5.0 / 72.0
    return h + 0.22


def slide_cards(prs, num, title, lead, cols):
    s = blank(prs)
    header(s, num, title, lead)
    top, bottom = BODY_TOP, BODY_BOTTOM
    n = len(cols)
    cw = (BODY_W - COL_GAP * (n - 1)) / n
    size = CANDS[-1]
    for cand in CANDS:
        if all(sum(card_h(h, b, cand, cw) for h, b in col) + CARD_GAP * (len(col) - 1) <= bottom - top for col in cols):
            size = cand
            break
    x = MARGIN
    overflow = False
    for col in cols:
        avail = (bottom - top) - CARD_GAP * (len(col) - 1)
        needs = [card_h(h, b, size, cw) for h, b in col]
        tot = sum(needs)
        if tot > avail:
            overflow = True
            heights = [nd * avail / tot for nd in needs]
        else:
            heights = [nd + (avail - tot) * (nd / tot) for nd in needs]
        y = top
        for (heading, bullets), h in zip(col, heights):
            rect(s, x, y, cw, h, fill=CARDBG, line_color=LINE)
            rect(s, x, y, cw, 0.05, fill=MID)
            hs = size + 2.5
            hl = est_lines(heading, hs, cw - 0.4)
            tb = textbox(s, x + 0.2, y + 0.11, cw - 0.4, hl * hs * 1.20 / 72.0 + 0.1)
            para(tb.text_frame, heading, size=hs, bold=True, color=DARK, first=True, line=1.1)
            tb = textbox(s, x + 0.2, y + 0.11 + hl * hs * 1.20 / 72.0 + 0.13, cw - 0.38, h - 0.4)
            tf = tb.text_frame
            for i, b in enumerate(bullets):
                strong = b.startswith("*")
                para(tf, b[1:] if strong else b, size=size, color=DARK if strong else INK, bold=strong,
                     first=(i == 0), space_after=5, bullet="▸", line=1.22)
            y += h + CARD_GAP
        x += cw + COL_GAP
    footer(s)
    print(f"slide {num:02d}: font {size}{'  OVERFLOW' if overflow else ''}")
    return s


def slide_title(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, fill=DARK)
    rect(s, 0, 5.55, SW, 1.95, fill=MID)
    rect(s, MARGIN, 0.55, 0.14, 1.05, fill=LIME)
    tb = textbox(s, 0.85, 0.50, 11.8, 1.1)
    para(tb.text_frame, "2026 커뮤니티 에너지 국제 컨퍼런스  |  한국에너지공단 · REN21", size=15, color=LIGHT, first=True)
    para(tb.text_frame, "주제 세션: 국제사회의 커뮤니티 에너지 사례 및 경험 — 국내 지자체 주도 커뮤니티 에너지", size=15, color=LIGHT)
    tb = textbox(s, MARGIN, 1.95, 12.2, 1.2)
    para(tb.text_frame, "지자체 주도 시민재생에너지 사업 추진 사례", size=38, bold=True, color=WHITE, first=True, line=1.05)
    tb = textbox(s, MARGIN, 3.10, 12.2, 2.3)
    para(tb.text_frame, "전남광주통합특별시 인수위원회 논의 — 「시민주권 공동체 자산 형성」 사업 구상", size=22, bold=True, color=LIME, first=True, line=1.1)
    para(tb.text_frame, "재생에너지 이익공유 → 마을월급 → 공동체 자산화", size=18, color=LIGHT, space_before=6)
    para(tb.text_frame, "Local-Government-Led Community Renewable Energy: The Jeonnam-Gwangju Transition Committee's "
         "Proposal on Citizen-Sovereign Community Wealth Building", size=15, color=LIGHT, space_before=8, line=1.15)
    tb = textbox(s, MARGIN, 5.75, 12.2, 1.6)
    para(tb.text_frame, "2026년 10월 1일(목)  한국프레스센터 국제회의장", size=16, bold=True, color=WHITE, first=True)
    para(tb.text_frame, "발표: 사회혁신교육원 사회적협동조합 이사장 김일영", size=16, color=WHITE, space_before=4)
    para(tb.text_frame, "※ 인수위원회 검토 단계의 구상을 소개하는 자료로, 조직체계와 예산은 미확정이라 다루지 않습니다.",
         size=15, color=LIGHT, space_before=8)
    return s


def slide_agenda(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, 1.12, fill=DARK)
    tb = textbox(s, MARGIN, 0.22, 12.2, 0.75)
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tb.text_frame, "발표 순서와 자료의 성격", size=24, bold=True, color=WHITE, first=True)
    items = [
        ("01–02", "왜 지자체가 나서는가", "추진 배경, 재생에너지를 출발점으로 삼는 이유"),
        ("03–04", "원리와 참고 모델", "자산기반 사회연대경제, 국내외 모델과 구양리"),
        ("05–07", "목표와 차별점", "사업 목표, 국가 햇빛소득마을과의 비교"),
        ("08–14", "4대 사업의 내용", "① 마을월급 기반 ② 기업 생태계 ③ 도메인 AI ④ 금융"),
        ("15–16", "생태계와 과제", "네트워크·규모화, 기대효과와 남은 과제"),
    ]
    y = 1.40
    for n, t, d in items:
        rect(s, MARGIN, y, 1.35, 0.92, fill=MID)
        tb = textbox(s, MARGIN, y + 0.24, 1.35, 0.5)
        para(tb.text_frame, n, size=16, bold=True, color=WHITE, first=True, align=PP_ALIGN.CENTER)
        rect(s, MARGIN + 1.35, y, 7.0, 0.92, fill=CARDBG, line_color=LINE)
        tb = textbox(s, MARGIN + 1.55, y + 0.07, 6.7, 0.85)
        para(tb.text_frame, t, size=17, bold=True, color=DARK, first=True, space_after=1)
        para(tb.text_frame, d, size=15, color=INK)
        y += 1.06
    rect(s, 9.15, 1.40, 3.63, 5.22, fill=LIGHT)
    rect(s, 9.15, 1.40, 3.63, 0.06, fill=AMBER)
    tb = textbox(s, 9.33, 1.58, 3.3, 5.0)
    tf = tb.text_frame
    para(tf, "이 자료의 성격", size=17, bold=True, color=DARK, first=True, space_after=8)
    for b in ["전남광주통합특별시 인수위원회의 검토용 사업계획 논의를 소개하는 발제",
              "확정 시책이 아닌 논의 단계의 구상. 조직체계·예산은 논의됐지만 미확정이라 제외",
              "'무엇을, 왜, 어떤 원리로' — 사업 내용과 설계 원리에 집중",
              "국가 햇빛소득마을과 경쟁이 아니라, 지자체가 그것을 확장·심화하는 사례"]:
        para(tf, b, size=15, color=INK, bullet="▸", space_after=7, line=1.2)
    return s


def slide_table(prs, num, title, lead, header_row, rows, widths, note=None):
    s = blank(prs)
    header(s, num, title, lead)
    top = 2.10
    bottom = 6.85 - (0.50 if note else 0)
    size = MIN
    def row_h(cells, sz):
        return max(est_lines(c, sz, w - 0.22) for c, w in zip(cells, widths)) * (sz * 1.25 / 72.0) + 0.18
    hh = 0.50
    for cand in CANDS:
        hs = [row_h(r, cand) for r in rows]
        if hh + sum(hs) <= bottom - top:
            size = cand
            break
    hs = [row_h(r, size) for r in rows]
    avail = bottom - top - hh
    tot = sum(hs)
    print(f"slide {num:02d} table: font {size}{'  OVERFLOW' if tot > avail else ''}")
    hs = [h * avail / tot for h in hs]
    x = MARGIN
    for c, w, col in zip(header_row, widths, (DARK, GRAY, MID)):
        rect(s, x, top, w, hh, fill=col, line_color=WHITE, line_w=1)
        tb = textbox(s, x, top + 0.09, w, hh)
        para(tb.text_frame, c, size=size + 1, bold=True, color=WHITE, first=True, align=PP_ALIGN.CENTER)
        x += w
    y = top + hh
    for ri, (r, h) in enumerate(zip(rows, hs)):
        x = MARGIN
        for ci, (c, w) in enumerate(zip(r, widths)):
            if ci == 0:
                fill, col, bold = LIGHT, DARK, True
            elif ci == 1:
                fill, col, bold = (WHITE if ri % 2 else CARDBG), INK, False
            else:
                fill, col, bold = RGBColor(0xEE, 0xF6, 0xF2) if ri % 2 else RGBColor(0xE4, 0xF1, 0xEA), DARK, False
            rect(s, x, y, w, h, fill=fill, line_color=WHITE, line_w=1)
            tb = textbox(s, x + 0.08, y + 0.04, w - 0.16, h - 0.06)
            tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            para(tb.text_frame, c, size=size, bold=bold, color=col, first=True, line=1.2,
                 align=PP_ALIGN.CENTER if ci == 0 else PP_ALIGN.LEFT)
            x += w
        y += h
    if note:
        rect(s, MARGIN, y + 0.08, BODY_W, 0.42, fill=RGBColor(0xFB, 0xF0, 0xE8))
        rect(s, MARGIN, y + 0.08, 0.08, 0.42, fill=AMBER)
        tb = textbox(s, MARGIN + 0.15, y + 0.12, BODY_W - 0.2, 0.38)
        para(tb.text_frame, note, size=MIN, bold=True, color=WARN, first=True)
    footer(s)
    return s


def slide_flow(prs, num, title, lead, boxes, formula, sub):
    s = blank(prs)
    header(s, num, title, lead)
    n = len(boxes)
    gap = 0.42
    bw = (BODY_W - gap * (n - 1)) / n
    top, bh = 2.15, 3.0
    x = MARGIN
    for i, (num_lbl, head, bullets) in enumerate(boxes):
        rect(s, x, top, bw, bh, fill=CARDBG, line_color=LINE)
        rect(s, x, top, bw, 0.95, fill=MID if i < 3 else DARK)
        tb = textbox(s, x + 0.12, top + 0.06, bw - 0.24, 0.88)
        tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tb.text_frame, f"{num_lbl} {head}", size=16, bold=True, color=WHITE, first=True, line=1.05)
        tb = textbox(s, x + 0.15, top + 1.05, bw - 0.3, bh - 1.1)
        tf = tb.text_frame
        for j, b in enumerate(bullets):
            para(tf, b, size=MIN, color=INK, first=(j == 0), bullet="▸", space_after=6, line=1.2)
        if i < n - 1:
            ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.05), Inches(top + bh / 2 - 0.2),
                                    Inches(gap - 0.1), Inches(0.4))
            ar.fill.solid(); ar.fill.fore_color.rgb = LIME; ar.line.fill.background(); ar.shadow.inherit = False
        x += bw + gap
    rect(s, MARGIN, top + bh + 0.12, BODY_W, 0.46, fill=LIGHT)
    tb = textbox(s, MARGIN, top + bh + 0.17, BODY_W, 0.4)
    para(tb.text_frame, "◀  ④의 수익·기금이 ①의 새 공동체와 ②의 신규 사업에 재투자되는 순환  ◀", size=MIN, bold=True,
         color=DARK, first=True, align=PP_ALIGN.CENTER)
    y = top + bh + 0.72
    rect(s, MARGIN, y, BODY_W, 0.62, fill=DARK)
    tb = textbox(s, MARGIN, y + 0.10, BODY_W, 0.5)
    para(tb.text_frame, formula, size=18, bold=True, color=LIME, first=True, align=PP_ALIGN.CENTER)
    tb = textbox(s, MARGIN, y + 0.66, BODY_W, 0.4)
    para(tb.text_frame, sub, size=MIN, color=GRAY, first=True, align=PP_ALIGN.CENTER)
    footer(s)
    return s


def slide_closing(prs, msg, sub, lines):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, fill=DARK)
    rect(s, MARGIN, 1.0, 0.14, 1.6, fill=LIME)
    tb = textbox(s, 0.9, 0.9, 11.9, 1.9)
    para(tb.text_frame, msg, size=30, bold=True, color=WHITE, first=True, line=1.15)
    para(tb.text_frame, sub, size=16, color=LIGHT, space_before=10, line=1.2)
    rect(s, MARGIN, 3.35, BODY_W, 0.03, fill=MID)
    tb = textbox(s, MARGIN, 3.55, BODY_W, 3.0)
    tf = tb.text_frame
    for i, l in enumerate(lines):
        para(tf, l, size=16, color=LIGHT, first=(i == 0), bullet="▸", space_after=14, line=1.3)
    tb = textbox(s, MARGIN, 6.75, BODY_W, 0.5)
    para(tb.text_frame, "감사합니다  |  Thank you", size=16, bold=True, color=LIME, first=True, align=PP_ALIGN.RIGHT)
    return s


# =====================================================================
prs = new_deck()
slide_title(prs)
slide_agenda(prs)

slide_cards(prs, 1, "추진 배경 — 왜 지자체가 재생에너지 사업에 나서는가",
    "AI 시대의 전력 수요 급증을 지역의 부담이 아닌, 시민의 자산과 소득으로 바꾸는 설계",
    [[("전력 수요의 시대 전환",
       ["AI 시대 전기 수요 급증으로 재생에너지 인프라 확충은 불가피",
        "*그 성과를 시민의 자산과 소득으로 환류하는 것이 핵심"]),
      ("전남광주의 조건",
       ["3대 메가프로젝트 핵심 지역, 재생에너지 잠재력이 큰 도농통합권",
        "재생에너지 생태계를 첫 단계로 도시·농어촌 생활경제 전반으로 확산"])],
     [("에너지 식민지인가, 주권자인가",
       ["집중·기업형 시스템은 지역이 생산 부담만 지고 이익은 역외 유출",
        "*분산형(지산지소형) 인프라로 지역이 주권자로서 공정한 혜택"]),
      ("두 광역의 자산을 승계·재설계",
       ["(광주) 사회적경제기업 1,300여 개소, 시민햇빛발전",
        "(전남) 해상풍력 등 거점, 햇빛·바람연금 등 주민참여형 이익공유",
        "*통합해 '도농상생형 시민주권경제'로 재설계"])]])

slide_cards(prs, 2, "왜 재생에너지로 공동체 자산을 형성하는가",
    "전기는 AI경제의 필수재, 그 생산은 가장 안정적인 사업 모델 — 공동체 자산 형성의 출발점",
    [[("안정성 — 투융자 리스크가 가장 낮다",
       ["전기 생산에서 재생에너지는 가장 안정적인 비즈니스 모델",
        "*시민기금·시민기업펀드의 투융자 리스크가 거의 없음"]),
      ("거버넌스형 사회연대경제의 토대",
       ["지역사회와 지방정부가 함께 만드는 퀘벡형 모델의 안정적 기반",
        "발전소 1개 사업이 아니라 지역이 함께 소유·운영하는 경제의 첫 단추"])],
     [("대자본형과 다른 경로",
       ["*주민공동체가 추진하고 국공유지 등 유휴지를 활용하는 것이 핵심",
        "총유(總有) 공동체 자산의 유지·관리가 지속성을 좌우"]),
      ("확장성 — 에너지에서 생활경제로",
       ["에너지전환에 따라 주거·지역산업·생활서비스로 확장",
        "*재생에너지 → 농어촌기본소득 → 사회연대경제로 기본사회 기반 확장"])]])

slide_cards(prs, 3, "설계 원리 — 자산기반 사회연대경제",
    "결핍이 아닌 지역의 자산에서 출발하고, 유출되던 부(富)를 공동체가 소유·순환시킨다",
    [[("이론적 배경 세 갈래",
       ["*지역자산 기반 발전론(ABCD) — 결핍이 아닌 지역 자산(사람·조직·자연·공간·재정)에서 출발",
        "*커뮤니티 웰스 빌딩 — 유출되던 부를 앵커기관–지역기업–주민 순환으로 축적. 소유·고용·조달·금융·토지 다섯 지렛대",
        "*사회연대경제·커먼즈 — 토지·에너지·데이터를 공동체가 민주적으로 소유·관리"])],
     [("다섯 가지 원칙",
       ["자산 소유의 민주화 · 호혜와 연대 · 지역순환",
        "민주적 거버넌스 · 자립과 회복력"]),
      ("일곱 가지 수단",
       ["공동체 토지신탁 · 협동조합복합체 · 공동체이익회사(CIC)",
        "참여예산 · 공동체기금 · 지역화폐 · 상호신용",
        "*재생에너지 사업은 이 수단들이 한 사업 안에서 결합하는 첫 무대"])]])

slide_cards(prs, 4, "참고한 국내외 모델 — 그리고 여주 구양리",
    "검증된 모델들을 전남광주의 도농 여건에 맞게 하나의 생태계로 통합하는 것이 이 구상의 위치",
    [[("해외 모델",
       ["몬드라곤 협동조합복합체 — 조합 간 연대와 내부 금융",
        "클리블랜드 에버그린 — 앵커기관 조달, 노동자 소유 기업",
        "프레스턴 모델 — 지역조달로 부의 역외 유출 차단",
        "영국 커뮤니티 주식 — 영구출자, 배당 제한, 1인 1표"]),
      ("국내 모델",
       ["신안 햇빛·바람연금 — 주민참여형 이익공유의 선례",
        "시민햇빛발전협동조합(광주), 완주 로컬푸드",
        "안산햇빛발전협동조합연합회 — 연합회형 모델"])],
     [("여주 구양리 — 마이크로그리드 결합형 마을",
       ["*1MW → 5MW 단계 고도화, 증설마다 마이크로그리드 단계 상향",
        "*5MW 도달 시 주민 1인당 월 50만원 배당 목표",
        "이번 컨퍼런스 현장견학지이자 전남광주 실증의 직접 참조 모델"]),
      ("전남광주 구상의 위치",
       ["*'재생에너지 이익공유 → 마을월급 → 공동체 자산화'로 발전시키는 통합 전략",
        "발전소 하나가 아니라 마을·읍면동·도시 조합을 잇는 생태계 설계"])]])

slide_cards(prs, 5, "사업 목표 — 지역 특성별 시민주권형 소득 체계",
    "주민과 시민에게 환류되는 소득 체계를 만들고, 그 기반이 되는 사회연대경제 생태계를 구축한다",
    [[("세 가지 목표",
       ["*① 주민과 시민에게 환류되는 지역 특성별 시민주권형 소득 체계",
        "*② 재생에너지·관광·가공 등 공동체 자산 기반 '마을월급' 소득활동",
        "*③ 다섯 원칙·일곱 수단에 기반한 자산기반 사회연대경제 생태계 인프라"]),
      ("도농상생형 시민주권경제로",
       ["광주의 사회적경제와 전남의 재생에너지·이익공유를 통합해 재설계",
        "마을(행정리) → 읍면동 → 도시 협동조합으로 대상 확대"])],
     [("정량 목표(검토안, 2026~2030 누계)",
       ["읍면동 자치공동체 250개소 · 마을 자치공동체 2,550개소",
        "우수 공동체이익회사 200개소 · 마을인프라 200개소",
        "시민기업 50개소 · 전문인력 교육 2,100명 · 활동 400명",
        "*인수위 검토 단계의 목표치이며 확정된 것이 아님"])]])

slide_table(prs, 6, "국가 햇빛소득마을과 무엇이 다른가 (1) 대상·금융·계통",
    "같은 계통·금융 조건 위에서 대상을 넓히고, 마을 소득 모델을 다각화한다",
    ["항목", "국가 햇빛소득마을", "전남광주형(검토안)"],
    [["대상", "행정리 단위의 마을", "행정리 마을 + 읍면 주민자치회 + 도시민 조합원 협동조합"],
     ["사업 내용", "햇빛발전소 설치에 의한 마을 소득에 한정", "동일 + 정원수·마을관광·농수산물가공 등 마을 특성 모델"],
     ["자부담·금융", "수익권 담보 시공비 융자와 이자 일부 지원, 주민보상금·지방소멸기금을 자부담 인정", "별도 시민기금으로 마을 자부담 저리 융자, 비정주 관계인구의 자부담 우선출자 인정"],
     ["계통·인프라", "계통 우선 연결, 불가 시 ESS 지원", "동일 + 계통 연결 불가 시 마이크로그리드 인프라 지원"]],
    [1.9, 4.75, 5.58])

slide_table(prs, 7, "국가 햇빛소득마을과 무엇이 다른가 (2) 운영·네트워크·지향",
    "발전소 설치(1단계)를 넘어 시민기업·경영공시·마이크로그리드·순환경제까지 설계",
    ["항목", "국가 햇빛소득마을", "전남광주형(검토안)"],
    [["시공·유지관리", "에너지공단 인증 레스코 담당", "시민출자·시 인증 시민기업이 레스코로 참여, 부품·마을 전기화 시민기업 육성"],
     ["네트워크", "선정 마을은 네트워크 의무 가입", "CIC 경영공시 의무가입, 연합회 가입으로 경영지원·생활서비스 사업서비스 이용"],
     ["지향", "재생에너지 발전사업에 한정(1단계)", "마이크로그리드로 에너지전환마을, 마을·지역순환경제를 처음부터 체계화"],
     ["지자체 역할", "이자 보전·교육컨설팅 한정, 자부담·계통연결 등 공무원 부담이 진입장벽", "광역·기초 원스톱 지원과 거버넌스 컨트롤타워, 확장모델의 국가 공모 제도화 견인"]],
    [1.9, 4.75, 5.58])

slide_flow(prs, 8, "사업 구조 — 서로 맞물린 4대 사업",
    "주민공동체가 자산을 발굴하고, 기업이 소득을 만들고, AI가 뒷받침하고, 금융이 환류·재투자한다",
    [("①", "마을월급 활동 기반 조성", ["주민자치회·마을회가 지역자산을 발굴해 소득모델 개발", "마을자산맵 · 국공유지 · 법인 · 초기 자본"]),
     ("②", "공동체이익회사·시민기업 생태계", ["CIC 창업·성장 지원, 마을경제 고도화 인프라", "가치사슬 허브 시민기업 — 연합회형·RESCO형"]),
     ("③", "AI 기반 마을월급 생태계", ["CIC·시민기업 업무자동화 도메인 특화 AI", "연대지능 AI 활동가와 필수 데이터 기여"]),
     ("④", "수익 환류·공동체 자산화", ["주민배당·기금 적립·재투자·복지로 환류", "신용보증 + 시민기금 재단 + 시민기업펀드"])],
    "재생에너지 이익공유  →  마을월급  →  공동체 자산화",
    "사업기간(검토안) 2026. 7 ~ 2030. 6  |  사업주체 전남광주통합특별시  |  세부 조직·예산은 미확정")

slide_cards(prs, 9, "① 마을월급 활동 기반 조성 — 주민공동체가 자산을 발굴한다",
    "발전소를 '주는' 사업이 아니라, 주민이 자기 자산을 조사·조직화하고 소득모델을 도출하는 과정을 지원",
    [[("누가 — 지원 대상",
       ["주민자치회·마을회 등 정주 주민공동체(관계인구 참여 컨소시엄 가점)",
        "*행정리 마을뿐 아니라 읍면동 자치공동체까지"]),
      ("①-1 지역자산 발굴·소득모델 개발",
       ["*마을자산맵 → 유휴 국공유지 사용권 → 법인 설립 → 초기 자본 → 주민활동 지원",
        "3단계: 자산조사·조직화 → 사용권·법인·자본 → 소득모델 도출"])],
     [("①-2 마을소득 다각화 컨설팅",
       ["재생에너지 / 마을관광 / 정원수(노지 스마트농업) / 농수산물 가공",
        "전문 컨설턴트 매칭 → 사업계획 고도화 → 공모 응모 → ② 진입"]),
      ("①-3 민주적 운영 역량과 과정기록",
       ["공동체 규약·수익배분 원칙·마을 비전·총회/이사회 훈련, 신탁계약 설계",
        "*합의문·회의록의 '과정기록'을 ②·④ 신청 필수 요건으로",
        "법인·금융·에너지·관광·농업·거버넌스 마을컨설턴트 풀"])]])

slide_cards(prs, 10, "② 공동체이익회사·시민기업 생태계 — 마을월급을 만드는 기업",
    "마을의 공동체이익회사(CIC)가 소득을 만들고, 광역·권역의 시민기업이 가치사슬을 잇는 두 층 구조",
    [[("공동체이익회사(CIC)란",
       ["지역사회 문제 해결과 공공이익을 목적으로 하는 사회연대경제형 기업",
        "마을 햇빛발전협동조합이 대표적 — 기술·경영·금융·청년인력 원스톱 지원"]),
      ("②-2 우수 CIC 성장지원 — 선정 기준",
       ["① 마을월급 창출 실적 ② 지역자산 기반성 ③ 사회연대경제 정체성",
        "④ 지속가능성 ⑤ 과정기록 보유"])],
     [("②-3 마을경제 고도화 인프라",
       ["마이크로그리드(100가구 기준) + 가공·정원수·관광 수익 인프라",
        "*요건: 마을 내 CIC, 마을회와 신탁계약(자산 장기 귀속), 수익배분 규약",
        "2026년 10여 개소 마이크로그리드 결합 실증부터 시작(에너지공단·한전 협력)"]),
      ("②-4 시민기업 — 가치사슬의 허브",
       ["발전협동조합 연합회형, 설치·운영·정산을 대행하는 RESCO형 등",
        "요건: 민주적 지배구조·이익 재투자·지역순환, 사회공헌 보고서, 경영공시"])]])

slide_cards(prs, 11, "지산지소 — 마이크로그리드와 마을 전기화의 원리",
    "지역에서 생산한 전기를 지역에서 쓰는 구조를 초기부터 설계해, 발전소 마을을 '에너지전환마을'로",
    [[("발전소 설치(1단계)에서 마을 전기화로",
       ["국가 사업이 발전소 설치 1단계라면, 전남광주는 마이크로그리드와 주거·마을경제 전기화",
        "*초기부터 마을 단위 전기화 전략 → 장기적으로 에너지전환마을"]),
      ("계통 문제를 다루는 순서",
       ["① 계통 우선접속을 국가 사업과 동일하게 확보",
        "② 용량 부족 시 ESS 지원  ③ 연결 불가 시 마이크로그리드 설치"])],
     [("단계적 고도화 — 구양리 방식",
       ["증설마다 마이크로그리드 단계를 상향해 자가소비·저장·관리 역량을 함께 키움",
        "실증 10여 개소 → 마이크로그리드 100개소 목표로 확산"]),
      ("분산형 인프라의 의미와 과제",
       ["*메가프로젝트 지역의 에너지 생산 부담을 줄이고, 지역이 주권자로서 혜택",
        "과제: 한전·VPP 기업·정부·지자체 간 설치비용 분담 합의 — 거버넌스가 선결"])]])

slide_cards(prs, 12, "③ AI 기반 마을월급 생태계 — 시민주권 특화 도메인 AI",
    "사회연대경제의 가치관과 현장의 암묵지를 학습한 특화 AI로 마을 기업과 활동가를 뒷받침",
    [[("왜 별도의 도메인 AI인가",
       ["일반 AI는 특정 분야에서 그럴듯한 오답을 내고 가치관 위험도 높음",
        "시민주권·사회연대경제·주민자치는 암묵지 영역이 많음",
        "*가치관과 암묵지를 학습한 시민주권 특화 AI를 독립 체계로 구축"]),
      ("③-1 마을월급 AI 플랫폼",
       ["CIC·시민기업 업무자동화(회계·행정·마케팅·공시) + 시민 정보·참여안내",
        "6단계 구축: 데이터 설계 → 축적 → 도메인 모델 → 개발 → 파일럿 → 운영"])],
     [("③-2 연대지능 AI 활동가",
       ["현장 활동가가 특화 AI를 훈련시킬 데이터 축적 과정의 기여자",
        "교육과정 수료 → 활동가 인증",
        "*필수 기여구조: 과정기록·사례 데이터의 플랫폼 축적이 활동·정산 요건"]),
      ("선순환의 설계",
       ["*활동가 데이터 → AI 고도화 → 기업 업무자동화 → 현장 확대 → 데이터 증가",
        "이력관리·경영정보 공개가 재정과 시민기금의 신뢰 기반"])]])

slide_cards(prs, 13, "④ 수익 환류와 공동체 자산화 — 사회연대경제 금융의 세 기둥",
    "지역자산의 수익을 주민배당·기금 적립·재투자·복지로 환류 — 기금 운영이 사업의 핵심 과제",
    [[("환류의 원리",
       ["지역자산 수익을 주민배당 · 시민기금 적립 · 재투자 · 복지로 환류",
        "*기금 운영이 핵심 과제 — 통합특별시가 지역금융기관과 협력체계 구축"]),
      ("[공공] 신용보증재단",
       ["CIC·시민기업 대출 보증으로 민간·금융기관 융자 연계",
        "보증재원의 레버리지 효과"])],
     [("[민간①] 시민기금 재단(크라우드펀딩)",
       ["기부형(선구매) / 융자형(저리 대출) / 투자형(평생지분투자)",
        "*투자형은 영국 커뮤니티 주식 참조 — 영구출자·배당 제한·1인 1표",
        "발전소 마을 자부담 부분에 한정하여 운영"]),
      ("[민간②] 시민기업펀드",
       ["성장단계 시민기업·CIC의 지분·메자닌 투자(스케일업 자본)",
        "임팩트투자 원칙, 배당금 회수로 재순환"])]])

slide_cards(prs, 14, "④ 수익 환류와 공동체 자산화 — 거버넌스와 공공의 역할",
    "초기에는 공공이 마중물을 대고, 시민참여가 커질수록 시민자본으로 대체한다 (규모·방식은 미확정)",
    [[("거버넌스와 투명성",
       ["시민기금: 총회 → 이사회(시민대표+전문가+공익) → 운용사무국",
        "투·융자 심사위의 사회적가치+재무 이중심사, 정보공시",
        "시민기업펀드: 출자자(LP)–운용자(GP) 구조, 투자심의위"]),
      ("자산의 장기 귀속 장치",
       ["공익신탁·시민펀딩·공공자산의 실질적 이전과 장기보유",
        "잉여자산의 호혜적 이용을 관리하는 공익기관 설립 지원"])],
     [("공공의 역할 — 마중물에서 시민자본으로",
       ["*초기 공공 출자 → 시민참여 확대에 따라 공공 지분 단계적 회수 → 시민자본 대체",
        "출연은 이자차액 보전·운영·플랫폼·교육 등 비회수 영역에 한정(경기도 사례 참조)",
        "로드맵: 초기(1~2년) 출자 → 중기(3~4년) 회수 개시 → 장기 시민자본 자립"])]])

slide_cards(prs, 15, "마을을 넘어 — 네트워크·연합회·규모화된 생활서비스",
    "마을들이 연대하는 생태계로 결속할 때, 소득은 '마을월급'을 넘어 지역경제가 된다",
    [[("전국 네트워크와 연합회",
       ["햇빛소득마을 선정 시 (햇빛배당전국)네트워크 의무 가입 → 연합회·전국 결속",
        "연합회형 예: 안산햇빛발전협동조합연합회",
        "*목표 1만 마을 × 100명 = 약 100만 명 규모"]),
      ("가치사슬 생태계의 상호거래 시장",
       ["CIC 연합회와 부품·시공·운영·유지관리 특화 조직이 상호거래 시장 구축",
        "*창업 지원 위주의 기존 정책과 달리 특정 산업의 가치사슬 생태계 조성"])],
     [("규모화된 생활서비스",
       ["마을 협동조합은 돌봄·문화·의료 생활서비스가 필요하지만 규모가 작아 유지 어려움",
        "*여러 마을을 규모화한 공급체계를 재생에너지 수익으로 뒷받침"]),
      ("투명성이 신뢰를 만든다",
       ["협동조합의 경영공시 의무 가입, 이력관리와 경영정보 공개",
        "재정과 시민기금의 신뢰를 위해 시작부터 체계적 정보 관리"])]])

slide_cards(prs, 16, "기대효과와 남은 과제",
    "발전소를 짓는 사업에서, 마을이 소유하고 지역이 순환시키는 경제로 — 거버넌스와 제도화는 아직 과제",
    [[("기대효과",
       ["*수익을 주민배당·기금 적립·재투자로 연결해 지속가능한 시민소득과 공동체 자산",
        "정주 주민과 관계인구의 협업으로 균형발전",
        "타 지역·대기업 의존도를 낮추는 지역순환경제",
        "수익 환원으로 시민 체감형 기본사회 기반 강화"])],
     [("남은 과제",
       ["*거버넌스: 한전·VPP·정부·지자체 간 마이크로그리드 비용 분담 합의",
        "*중앙정부 협력: 에너지공단 MOU·TF, 선도 모델 연구, 확장모델의 국가 공모 제도화"]),
      ("확정을 기다리는 것",
       ["부서 간 중복을 가로지르는 통합 조정 체계",
        "전담 조직과 예산·재원 조달 방식 — 논의됐지만 미확정이라 제외"])]])

slide_closing(prs,
    "발전소를 짓는 사업에서,\n마을이 소유하고 지역이 순환시키는 경제로",
    "전남광주통합특별시 인수위원회 「시민주권 공동체 자산 형성」 논의 — 재생에너지 이익공유 → 마을월급 → 공동체 자산화",
    ["재생에너지는 가장 안정적인 사업 모델이므로 공동체 자산 형성의 출발점으로 가장 적합하다",
     "주민공동체가 자산을 발굴하고(①), 기업이 소득을 만들고(②), 도메인 AI가 뒷받침하며(③), 금융이 환류·재투자한다(④)",
     "국가 햇빛소득마을을 읍면동형·도시형, 마이크로그리드·마을 전기화, 순환경제로 확장하는 지자체 모델",
     "조직체계와 예산은 논의됐지만 미확정 — 사업 내용과 설계 원리를 소개하는 검토 단계의 구상입니다"])

prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))
