# -*- coding: utf-8 -*-
"""English edition of the Jeonnam-Gwangju conference deck (body text >= 15pt), mirroring the submitted Korean PDF."""
import sys, math
sys.path.insert(0, r"C:\Users\User\.claude\skills\pptx-lecture-deck\scripts")
import deck_lib
deck_lib.FONT = "Calibri"
from deck_lib import *
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

OUT = sys.argv[1]
AMBER = RGBColor(0xE0, 0x9F, 0x3E)
MIN = 15.0
CANDS = (17.0, 16.0, 15.0)
CHAR_W = 0.55   # average width of an English character in em (Korean estimator uses 0.88)


def est_lines(text, size, width_in):
    char_w = size * CHAR_W / 72.0
    per_line = max(10, int(width_in / char_w))
    return max(1, math.ceil(len(text) / per_line))


def header(s, num, title, lead):
    rect(s, 0, 0, SW, 1.12, fill=DARK)
    rect(s, MARGIN, 0.24, 0.66, 0.64, fill=LIME)
    tb = textbox(s, MARGIN, 0.30, 0.66, 0.56)
    para(tb.text_frame, f"{num:02d}", size=17, bold=True, color=DARK, first=True, align=PP_ALIGN.CENTER)
    tb = textbox(s, 1.42, 0.14, 11.4, 0.92)
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tb.text_frame, title, size=22 if len(title) > 62 else 24, bold=True, color=WHITE, first=True, line=1.0)
    rect(s, 0, 1.12, SW, 0.78, fill=LIGHT)
    rect(s, 0, 1.12, 0.12, 0.78, fill=MID)
    tb = textbox(s, MARGIN, 1.14, 12.3, 0.74)
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tb.text_frame, lead, size=16 if len(lead) < 105 else 15, bold=True, color=DARK, first=True, line=1.08)


def card_h(heading, bullets, size, cw):
    tw = cw - 0.42
    hs = size + 2.5
    h = 0.16 + est_lines(heading, hs, cw - 0.4) * (hs * 1.20 / 72.0) + 0.15
    for b in bullets:
        t = b[1:] if b.startswith("*") else b
        h += est_lines("> " + t, size, tw) * (size * 1.22 / 72.0) + 5.0 / 72.0
    return h + 0.22


def slide_cards(prs, num, title, lead, cols):
    s = blank(prs)
    header(s, num, title, lead)
    top, bottom = BODY_TOP, 7.08
    n = len(cols)
    cw = (BODY_W - COL_GAP * (n - 1)) / n
    size = CANDS[-1]
    for cand in CANDS:
        if all(sum(card_h(h, b, cand, cw) for h, b in col) + CARD_GAP * (len(col) - 1) <= bottom - top for col in cols):
            size = cand
            break
    x = MARGIN
    overflow = False
    for col in cols:
        avail = (bottom - top) - CARD_GAP * (len(col) - 1)
        needs = [card_h(h, b, size, cw) for h, b in col]
        tot = sum(needs)
        if tot > avail:
            overflow = True
            heights = [nd * avail / tot for nd in needs]
        else:
            heights = [nd + (avail - tot) * (nd / tot) for nd in needs]
        y = top
        for (heading, bullets), h in zip(col, heights):
            rect(s, x, y, cw, h, fill=CARDBG, line_color=LINE)
            rect(s, x, y, cw, 0.05, fill=MID)
            hs = size + 2.5
            hl = est_lines(heading, hs, cw - 0.4)
            tb = textbox(s, x + 0.2, y + 0.11, cw - 0.4, hl * hs * 1.20 / 72.0 + 0.1)
            para(tb.text_frame, heading, size=hs, bold=True, color=DARK, first=True, line=1.1)
            tb = textbox(s, x + 0.2, y + 0.11 + hl * hs * 1.20 / 72.0 + 0.13, cw - 0.38, h - 0.4)
            tf = tb.text_frame
            for i, b in enumerate(bullets):
                strong = b.startswith("*")
                para(tf, b[1:] if strong else b, size=size, color=DARK if strong else INK, bold=strong,
                     first=(i == 0), space_after=5, bullet="▸", line=1.2)
            y += h + CARD_GAP
        x += cw + COL_GAP
    print(f"slide {num:02d}: font {size}{'  OVERFLOW' if overflow else ''}")
    return s


def slide_title(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, fill=DARK)
    rect(s, 0, 5.55, SW, 1.95, fill=MID)
    rect(s, MARGIN, 0.55, 0.14, 1.05, fill=LIME)
    tb = textbox(s, 0.85, 0.50, 11.8, 1.1)
    para(tb.text_frame, "2026 Community Energy International Conference  |  Korea Energy Agency · REN21", size=15, color=LIGHT, first=True)
    para(tb.text_frame, "Session: Community Energy Cases and Experiences Worldwide — Local Government-Led Community Energy in Korea", size=15, color=LIGHT)
    tb = textbox(s, MARGIN, 1.85, 12.2, 1.3)
    para(tb.text_frame, "Local Government-Led Citizen Renewable Energy: A Case from Korea", size=34, bold=True, color=WHITE, first=True, line=1.05)
    tb = textbox(s, MARGIN, 3.05, 12.2, 2.4)
    para(tb.text_frame, "The Jeonnam-Gwangju Integrated Special City Transition Committee's Deliberations on the "
         "\"Citizen-Sovereign Community Wealth Building\" Initiative", size=20, bold=True, color=LIME, first=True, line=1.1)
    para(tb.text_frame, "Renewable energy benefit sharing → Village Salary → Community wealth", size=18, color=LIGHT, space_before=6)
    para(tb.text_frame, "지자체 주도 시민재생에너지 사업 추진 사례 — 전남광주통합특별시 인수위원회 논의 「시민주권 공동체 자산 형성」",
         size=15, color=LIGHT, space_before=8)
    tb = textbox(s, MARGIN, 5.75, 12.2, 1.6)
    para(tb.text_frame, "Thursday, 1 October 2026  |  Korea Press Center, International Conference Hall, Seoul", size=16, bold=True, color=WHITE, first=True)
    para(tb.text_frame, "Presenter: Il-young Kim, Chairperson, Social Innovation Platform", size=16, color=WHITE, space_before=4)
    para(tb.text_frame, "※ This material introduces a proposal under review by the Transition Committee. Organizational structure and budget are not yet finalized and are not covered.",
         size=15, color=LIGHT, space_before=8)
    return s


def slide_agenda(prs):
    s = blank(prs)
    rect(s, 0, 0, SW, 1.12, fill=DARK)
    tb = textbox(s, MARGIN, 0.22, 12.2, 0.75)
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tb.text_frame, "Outline and Nature of This Material", size=24, bold=True, color=WHITE, first=True)
    items = [
        ("01–02", "Why a local government steps in", "Background; why renewable energy is the starting point"),
        ("03–04", "Principles and reference models", "Asset-based SSE; models at home and abroad; Guyang-ri"),
        ("05–07", "Objectives and what is different", "Objectives; comparison with the national Sunlight Income Village program"),
        ("08–14", "The four programs", "① Village Salary base ② enterprises ③ domain AI ④ finance"),
        ("15–16", "Ecosystem and remaining tasks", "Networks and scale; expected effects and open issues"),
    ]
    y = 1.40
    for n, t, d in items:
        rect(s, MARGIN, y, 1.35, 0.92, fill=MID)
        tb = textbox(s, MARGIN, y + 0.24, 1.35, 0.5)
        para(tb.text_frame, n, size=16, bold=True, color=WHITE, first=True, align=PP_ALIGN.CENTER)
        rect(s, MARGIN + 1.35, y, 7.0, 0.92, fill=CARDBG, line_color=LINE)
        tb = textbox(s, MARGIN + 1.55, y + 0.07, 6.7, 0.85)
        para(tb.text_frame, t, size=17, bold=True, color=DARK, first=True, space_after=1)
        para(tb.text_frame, d, size=15, color=INK)
        y += 1.06
    rect(s, 9.15, 1.40, 3.63, 5.22, fill=LIGHT)
    rect(s, 9.15, 1.40, 3.63, 0.06, fill=AMBER)
    tb = textbox(s, 9.33, 1.58, 3.3, 5.0)
    tf = tb.text_frame
    para(tf, "About this material", size=17, bold=True, color=DARK, first=True, space_after=8)
    for b in ["Introduces the deliberations of the Jeonnam-Gwangju Integrated Special City Transition Committee on a draft plan under review",
              "A proposal under discussion, not an adopted policy. Organization and budget were discussed but remain undecided and are excluded",
              "Focus on 'what, why, and by what principles' — program content and design principles",
              "Not a rival to the national Sunlight Income Village program, but a case of how a local government can extend it"]:
        para(tf, b, size=15, color=INK, bullet="▸", space_after=7, line=1.18)
    return s


def slide_table(prs, num, title, lead, header_row, rows, widths):
    s = blank(prs)
    header(s, num, title, lead)
    top, bottom = 2.10, 7.05
    size = MIN
    def row_h(cells, sz):
        return max(est_lines(c, sz, w - 0.22) for c, w in zip(cells, widths)) * (sz * 1.25 / 72.0) + 0.18
    hh = 0.50
    for cand in CANDS:
        hs = [row_h(r, cand) for r in rows]
        if hh + sum(hs) <= bottom - top:
            size = cand
            break
    hs = [row_h(r, size) for r in rows]
    avail = bottom - top - hh
    tot = sum(hs)
    print(f"slide {num:02d} table: font {size}{'  OVERFLOW' if tot > avail else ''}")
    hs = [h * avail / tot for h in hs]
    x = MARGIN
    for c, w, col in zip(header_row, widths, (DARK, GRAY, MID)):
        rect(s, x, top, w, hh, fill=col, line_color=WHITE, line_w=1)
        tb = textbox(s, x, top + 0.09, w, hh)
        para(tb.text_frame, c, size=size + 1, bold=True, color=WHITE, first=True, align=PP_ALIGN.CENTER)
        x += w
    y = top + hh
    for ri, (r, h) in enumerate(zip(rows, hs)):
        x = MARGIN
        for ci, (c, w) in enumerate(zip(r, widths)):
            if ci == 0:
                fill, col, bold = LIGHT, DARK, True
            elif ci == 1:
                fill, col, bold = (WHITE if ri % 2 else CARDBG), INK, False
            else:
                fill, col, bold = RGBColor(0xEE, 0xF6, 0xF2) if ri % 2 else RGBColor(0xE4, 0xF1, 0xEA), DARK, False
            rect(s, x, y, w, h, fill=fill, line_color=WHITE, line_w=1)
            tb = textbox(s, x + 0.08, y + 0.04, w - 0.16, h - 0.06)
            tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            para(tb.text_frame, c, size=size, bold=bold, color=col, first=True, line=1.18,
                 align=PP_ALIGN.CENTER if ci == 0 else PP_ALIGN.LEFT)
            x += w
        y += h
    return s


def slide_flow(prs, num, title, lead, boxes, formula, sub):
    s = blank(prs)
    header(s, num, title, lead)
    n = len(boxes)
    gap = 0.42
    bw = (BODY_W - gap * (n - 1)) / n
    top, bh = 2.15, 3.1
    x = MARGIN
    for i, (num_lbl, head, bullets) in enumerate(boxes):
        rect(s, x, top, bw, bh, fill=CARDBG, line_color=LINE)
        rect(s, x, top, bw, 0.95, fill=MID if i < 3 else DARK)
        tb = textbox(s, x + 0.12, top + 0.06, bw - 0.24, 0.88)
        tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tb.text_frame, f"{num_lbl} {head}", size=16, bold=True, color=WHITE, first=True, line=1.05)
        tb = textbox(s, x + 0.15, top + 1.05, bw - 0.3, bh - 1.1)
        tf = tb.text_frame
        for j, b in enumerate(bullets):
            para(tf, b, size=MIN, color=INK, first=(j == 0), bullet="▸", space_after=6, line=1.18)
        if i < n - 1:
            ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.05), Inches(top + bh / 2 - 0.2),
                                    Inches(gap - 0.1), Inches(0.4))
            ar.fill.solid(); ar.fill.fore_color.rgb = LIME; ar.line.fill.background(); ar.shadow.inherit = False
        x += bw + gap
    rect(s, MARGIN, top + bh + 0.12, BODY_W, 0.46, fill=LIGHT)
    tb = textbox(s, MARGIN, top + bh + 0.17, BODY_W, 0.4)
    para(tb.text_frame, "◀  Returns and funds from ④ are reinvested in new communities (①) and new projects (②)  ◀", size=MIN, bold=True,
         color=DARK, first=True, align=PP_ALIGN.CENTER)
    y = top + bh + 0.72
    rect(s, MARGIN, y, BODY_W, 0.62, fill=DARK)
    tb = textbox(s, MARGIN, y + 0.10, BODY_W, 0.5)
    para(tb.text_frame, formula, size=18, bold=True, color=LIME, first=True, align=PP_ALIGN.CENTER)
    tb = textbox(s, MARGIN, y + 0.66, BODY_W, 0.4)
    para(tb.text_frame, sub, size=MIN, color=GRAY, first=True, align=PP_ALIGN.CENTER)
    return s


def slide_closing(prs, msg, sub, lines):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, fill=DARK)
    rect(s, MARGIN, 0.95, 0.14, 1.6, fill=LIME)
    tb = textbox(s, 0.9, 0.85, 11.9, 2.1)
    para(tb.text_frame, msg, size=28, bold=True, color=WHITE, first=True, line=1.15)
    para(tb.text_frame, sub, size=16, color=LIGHT, space_before=10, line=1.2)
    rect(s, MARGIN, 3.45, BODY_W, 0.03, fill=MID)
    tb = textbox(s, MARGIN, 3.65, BODY_W, 3.0)
    tf = tb.text_frame
    for i, l in enumerate(lines):
        para(tf, l, size=16, color=LIGHT, first=(i == 0), bullet="▸", space_after=12, line=1.25)
    tb = textbox(s, MARGIN, 6.75, BODY_W, 0.5)
    para(tb.text_frame, "Thank you  |  감사합니다", size=16, bold=True, color=LIME, first=True, align=PP_ALIGN.RIGHT)
    return s


# =====================================================================
prs = new_deck()
slide_title(prs)
slide_agenda(prs)

slide_cards(prs, 1, "Background — Why Does a Local Government Take On Renewable Energy?",
    "Turning the AI-era surge in power demand into citizens' assets and income, not a burden on the region",
    [[("A new era of power demand",
       ["Surging AI-era power demand makes renewable expansion unavoidable",
        "*The key is to channel the gains into citizens' assets and income"]),
      ("Jeonnam-Gwangju's conditions",
       ["Core region of three national mega-projects; urban–rural area with large renewable potential",
        "Build the renewable ecosystem first, then spread to everyday urban and rural economies"])],
     [("Energy colony, or energy sovereign?",
       ["Centralized corporate systems leave the region the burden while profits flow out",
        "*Distributed local-production, local-consumption infrastructure gives the region a fair share"]),
      ("Inheriting two regions' assets",
       ["(Gwangju) ~1,300 social economy enterprises; citizens' solar cooperatives",
        "(Jeonnam) offshore wind hubs; Sunlight and Wind Pension benefit sharing",
        "*Combined into an 'urban–rural, citizen-sovereign economy'"])]])

slide_cards(prs, 2, "Why Build Community Wealth with Renewable Energy?",
    "Electricity is the essential good of the AI economy, and generating it is the most stable business — the starting point for community wealth",
    [[("Stability — the lowest investment risk",
       ["Among electricity businesses, renewables are the most stable model",
        "*Investment and lending risk for the Citizen Fund and the Citizen Enterprise Fund is close to nil"]),
      ("A base for a governance-type SSE",
       ["A stable foundation for a Québec-style model built jointly by the community and local government",
        "Not a single power plant, but the first step toward an economy the region owns and runs together"])],
     [("A different path from big capital",
       ["*Resident communities lead the projects, using idle public land and other unused sites",
        "Maintaining collectively owned (chongyu) community assets decides sustainability"]),
      ("Scalability — from energy to the everyday economy",
       ["Expands with the energy transition into housing, local industry and living services",
        "*Renewables → rural basic income → social and solidarity economy: a foundation for a 'basic society'"])]])

slide_cards(prs, 3, "Design Principles — Asset-Based Social and Solidarity Economy",
    "Start from the region's assets, not its deficits, and let the community own and circulate the wealth that used to leak out",
    [[("Three theoretical roots",
       ["*Asset-Based Community Development (ABCD) — start from local assets (people, organizations, nature, space, finance), not deficits",
        "*Community Wealth Building — keep leaking wealth local through anchor institution–local business–resident circulation; five levers: ownership, employment, procurement, finance, land",
        "*Social and solidarity economy and the commons — land, energy and data democratically owned and managed by the community"])],
     [("Five principles",
       ["Democratized asset ownership · reciprocity and solidarity · local circulation",
        "Democratic governance · self-reliance and resilience"]),
      ("Seven instruments",
       ["Community land trust · cooperative complex · Community Interest Company (CIC)",
        "Participatory budgeting · community fund · local currency · mutual credit",
        "*Renewable energy is the first arena where these instruments combine in a single project"])]])

slide_cards(prs, 4, "Reference Models at Home and Abroad — and Guyang-ri, Yeoju",
    "This proposal's role is to integrate proven models into one ecosystem fitted to Jeonnam-Gwangju's urban–rural conditions",
    [[("International models",
       ["Mondragón — inter-cooperative solidarity, internal finance",
        "Cleveland Evergreen — anchor procurement, worker-owned firms",
        "Preston Model — local procurement keeps wealth local",
        "UK community shares — withdrawable capital, capped interest, one member one vote"]),
      ("Domestic models",
       ["Shinan Sunlight and Wind Pension — resident benefit sharing",
        "Citizens' Solar Cooperative (Gwangju); Wanju Local Food",
        "Ansan Federation of Solar Cooperatives — federation model"])],
     [("Guyang-ri, Yeoju — a microgrid-linked village",
       ["*Staged upgrade 1 MW → 5 MW; microgrid level rises with each expansion",
        "*Target: KRW 500,000 per resident per month at 5 MW",
        "This conference's site-visit village; direct reference for Jeonnam-Gwangju pilots"]),
      ("Where Jeonnam-Gwangju's proposal stands",
       ["*Integrated strategy: 'renewable benefit sharing → Village Salary → community wealth'",
        "Not one plant, but an ecosystem linking villages, townships and urban cooperatives"])]])

slide_cards(prs, 5, "Objectives — A Citizen-Sovereign Income System Tailored to Each Locality",
    "Build an income system that flows back to residents and citizens, and the social and solidarity economy ecosystem beneath it",
    [[("Three objectives",
       ["*① A locally tailored, citizen-sovereign income system for residents and citizens",
        "*② 'Village Salary' income from community assets — renewables, tourism, processing",
        "*③ An asset-based SSE ecosystem built on five principles and seven instruments"]),
      ("Toward an urban–rural mutual prosperity economy",
       ["Combine Gwangju's social economy with Jeonnam's renewables and benefit sharing",
        "Widen scope: villages (ri) → townships (eup/myeon/dong) → urban cooperatives"])],
     [("Quantitative targets (draft, cumulative 2026–2030)",
       ["250 township and 2,550 village community organizations",
        "200 leading CICs · 200 village infrastructure sites · 50 citizen enterprises",
        "2,100 trained specialists · 400 active practitioners",
        "*Draft figures under review, not yet adopted"])]])

slide_table(prs, 6, "How Is It Different from the National Sunlight Income Village? (1) Scope, Finance, Grid",
    "Same grid and finance conditions, but a wider scope and diversified village income models",
    ["Item", "National Sunlight Income Village", "Jeonnam-Gwangju model (draft)"],
    [["Scope", "Administrative villages (ri)", "Villages + township Residents' Autonomy Councils + cooperatives with urban members"],
     ["Business", "Village income limited to a solar power plant", "Same + locally tailored models: landscape trees, village tourism, agri-food processing"],
     ["Own contribution and finance", "Construction loan secured by revenue rights with partial interest support; resident compensation and the Local Depopulation Fund count as own contribution", "Low-interest loans for the village share from a separate Citizen Fund; priority equity from non-resident stakeholders counts as own contribution"],
     ["Grid and infrastructure", "Priority grid connection; ESS support where connection is not possible", "Same + microgrid infrastructure where grid connection is not possible"]],
    [1.9, 4.75, 5.58])

slide_table(prs, 7, "How Is It Different from the National Sunlight Income Village? (2) Operation, Network, Direction",
    "Beyond plant installation (stage 1): citizen enterprises, management disclosure, microgrids and a circular economy",
    ["Item", "National Sunlight Income Village", "Jeonnam-Gwangju model (draft)"],
    [["Construction and O&M", "KEA-certified RESCOs (renewable energy service companies)", "Citizen-invested, city-certified citizen enterprises act as RESCOs; fostering component, construction and village-electrification enterprises"],
     ["Network", "Selected villages must join the national network", "CICs must join the management disclosure system; federation membership gives access to management support and living-service business services"],
     ["Direction", "Limited to renewable generation (stage 1)", "Energy transition villages via microgrids; village and regional circular economy systematized from the start"],
     ["Local government role", "Limited to interest subsidy and training; own contribution and grid connection burden officials and create entry barriers", "One-stop support at metropolitan and municipal level with a governance control tower; drive institutionalization of the extended model in national calls"]],
    [1.9, 4.75, 5.58])

slide_flow(prs, 8, "Program Structure — Four Interlocking Programs",
    "Residents discover assets, enterprises create income, AI supports them, and finance recycles and reinvests the returns",
    [("①", "Village Salary foundations", ["Residents' councils and village associations discover assets and develop income models", "Asset map · public land · entity · seed capital"]),
     ("②", "CIC and citizen enterprise ecosystem", ["CIC start-up and growth support; village infrastructure", "Value-chain hub citizen enterprises (federation, RESCO)"]),
     ("③", "AI-based Village Salary ecosystem", ["Domain AI automating CIC and citizen enterprise operations", "Solidarity Intelligence AI Activists; mandatory data contribution"]),
     ("④", "Returns and community wealth", ["Dividends, fund reserves, reinvestment, welfare", "Credit guarantee + Citizen Fund + Enterprise Fund"])],
    "Renewable energy benefit sharing  →  Village Salary  →  Community wealth",
    "Period (draft) July 2026 – June 2030  |  Lead: Jeonnam-Gwangju Integrated Special City  |  Organization and budget not yet finalized")

slide_cards(prs, 9, "① Village Salary Foundations — Residents Discover Their Own Assets",
    "Not a program that 'hands out' a power plant, but support for residents to survey and organize their assets and derive an income model",
    [[("Who — eligible communities",
       ["Residents' Autonomy Councils, village associations and similar (bonus for consortia with non-resident stakeholders)",
        "*Also township community organizations and city/county/district citizen communities"]),
      ("①-1 Asset discovery and income models",
       ["*Village asset map → idle public land → legal entity → seed capital → activity support",
        "Three stages: survey and organizing → rights, entity, capital → income model"])],
     [("①-2 Income diversification consulting",
       ["Renewables / village tourism / landscape trees (open-field smart farming) / agri-food processing",
        "Consultant matching → business plan upgrade → funding calls → entry into ②"]),
      ("①-3 Democratic capacity and process records",
       ["Training in bylaws, profit-sharing rules, village vision, meetings and boards; trust agreements",
        "*'Process records' (agreements, minutes) required for applying to ② and ④",
        "Village consultant pool: legal, finance, energy, tourism, agriculture, governance"])]])

slide_cards(prs, 10, "② CIC and Citizen Enterprise Ecosystem — The Enterprises That Create the Village Salary",
    "A two-layer structure: village CICs create income, and metropolitan and regional citizen enterprises link the value chain",
    [[("What is a Community Interest Company (CIC)?",
       ["An SSE enterprise set up to solve local problems and serve the public interest",
        "Village solar cooperatives are typical — one-stop support in technology, management, finance, staffing"]),
      ("②-2 Leading CIC growth support — criteria",
       ["① Village Salary generated ② local asset base ③ SSE identity",
        "④ sustainability ⑤ process records on file"])],
     [("②-3 Village economy infrastructure",
       ["Microgrids (per 100 households) + processing, landscape-tree and tourism infrastructure",
        "*Requirements: a village CIC, a trust agreement with the village association (asset lock), profit-sharing bylaws",
        "Microgrid-linked pilots from 2027 (with KEA and KEPCO)"]),
      ("②-4 Leading citizen enterprises — value-chain hubs",
       ["Federations of solar cooperatives; RESCO-type firms that install, operate and settle for villages",
        "Requirements: democratic governance, profit reinvestment, local circulation, disclosure"])]])

slide_cards(prs, 11, "Local Production, Local Consumption — The Logic of Microgrids and Village Electrification",
    "Design from the outset a structure in which locally generated power is used locally, turning power-plant villages into 'energy transition villages'",
    [[("From plant installation (stage 1) to village electrification",
       ["The national program installs plants; Jeonnam-Gwangju adds microgrids and electrifies housing and the village economy",
        "*Village-level electrification from the start → energy transition villages"]),
      ("The order for handling grid issues",
       ["① Secure priority grid connection on the same terms as the national program",
        "② ESS support where capacity is short  ③ microgrids where connection is impossible"])],
     [("Staged upgrading — the Guyang-ri way",
       ["Raise the microgrid level with each capacity expansion, building self-consumption, storage and management capacity together",
        "Some 10 pilot sites → scale up toward 100 microgrids"]),
      ("Meaning and challenge of distributed infrastructure",
       ["*Ease the generation burden on mega-project regions; the region benefits as sovereign",
        "Challenge: cost sharing among KEPCO, VPP operators, national and local government — governance first"])]])

slide_cards(prs, 12, "③ AI-Based Village Salary Ecosystem — Domain AI for Citizen Sovereignty",
    "Specialized AI trained on the values and tacit knowledge of the social and solidarity economy supports village enterprises and practitioners",
    [[("Why a separate domain AI?",
       ["General-purpose AI gives plausible wrong answers in specialized fields and carries value risks",
        "Citizen sovereignty, the social and solidarity economy and residents' autonomy rest heavily on tacit knowledge",
        "*Build an independent citizen-sovereignty AI trained on those values and that tacit knowledge"]),
      ("③-1 Village Salary AI platform",
       ["Automation for CICs and citizen enterprises (accounting, administration, marketing, disclosure) + information and participation guidance for citizens",
        "Six-step build: data design → accumulation → domain model → development → pilot → operation"])],
     [("③-2 Solidarity Intelligence AI Activists",
       ["Field practitioners who contribute the data that trains the specialized AI",
        "Course completion → activist certification",
        "*Mandatory contribution: uploading process records and case data to the platform is a condition of activity and payment"]),
      ("Designing the virtuous cycle",
       ["*Practitioner data → better AI → enterprise automation → more field activity → more data",
        "Track records and management disclosure underpin trust in public finance and the Citizen Fund"])]])

slide_cards(prs, 13, "④ Returns and Community Wealth — Three Pillars of Social and Solidarity Finance",
    "Channel returns from local assets into resident dividends, fund reserves, reinvestment and welfare — fund management is the core task",
    [[("The logic of recirculation",
       ["Returns flow to resident dividends · Citizen Fund reserves · reinvestment · welfare",
        "*Fund management is the core task — partnership with local financial institutions"]),
      ("[Public] Credit Guarantee Foundation",
       ["Loan guarantees for CICs and citizen enterprises unlock private and bank lending",
        "Leverage on the guarantee reserve"])],
     [("[Private ①] Citizen Fund Foundation (crowdfunding)",
       ["Donation type (pre-purchase) / loan type (low interest) / investment type (lifetime equity)",
        "*Follows UK community shares — non-transferable, no capital gain, withdrawal at par at the society's discretion",
        "Used only for villages' own share of power plants"]),
      ("[Private ②] Citizen Enterprise Fund",
       ["Equity and mezzanine investment in growth-stage citizen enterprises and CICs",
        "Impact investment principles; dividends recovered and recycled"])]])

slide_cards(prs, 14, "④ Returns and Community Wealth — Governance and the Public Role",
    "Public money primes the pump at first and is replaced by citizen capital as participation grows (scale and method not yet decided)",
    [[("Governance and transparency",
       ["Citizen Fund: general meeting → board (citizen representatives + experts + public interest) → management office",
        "The investment and loan committee applies a dual social-value and financial review; information disclosure",
        "Citizen Enterprise Fund: LP–GP structure with an investment review committee"]),
      ("Locking assets in for the long term",
       ["Charitable trusts, citizen funding, real transfer and long-term holding of public assets",
        "Support for public-interest institutions that manage the reciprocal use of surplus assets"])],
     [("The public role — from seed money to citizen capital",
       ["*Initial public investment → staged recovery of the public stake as citizen participation grows → replaced by citizen capital",
        "Grants limited to non-recoverable items: interest-gap subsidy, operations, platform, education (cf. Gyeonggi Province)",
        "Roadmap: years 1–2 investment → years 3–4 start of recovery → long-term citizen-capital self-reliance"])]])

slide_cards(prs, 15, "Beyond the Village — Networks, Federations and Living Services at Scale",
    "When villages bond into a solidarity ecosystem, income grows from a 'Village Salary' into a regional economy",
    [[("National network and federations",
       ["Selected villages must join the national Sunlight Dividend Network → federation and nationwide bonds",
        "Example: Ansan Federation of Solar Power Cooperatives",
        "*Target: 10,000 villages × 100 members ≈ one million people"]),
      ("A mutual-trading market across the value chain",
       ["CIC federations and specialist firms in components, construction and O&M build a mutual-trading market",
        "*Unlike start-up-centered policy, this builds a value-chain ecosystem for one industry"])],
     [("Living services at scale",
       ["Village cooperatives need care, culture and health services but are too small to sustain them alone",
        "*A supply system pooled across villages, underwritten by renewable revenues"]),
      ("Transparency builds trust",
       ["Mandatory management disclosure for cooperatives; track records and management information made public",
        "Systematic information management from day one to earn trust in public finance and the Citizen Fund"])]])

slide_cards(prs, 16, "Expected Effects and Remaining Tasks",
    "From building power plants to an economy villages own and the region circulates — governance and institutionalization remain open",
    [[("Expected effects",
       ["*Sustainable citizen income and community wealth by linking returns to dividends, reserves and reinvestment",
        "Balanced development through cooperation between residents and non-resident stakeholders",
        "A circular regional economy less dependent on outside regions and large corporations",
        "A tangible 'basic society' foundation through returns to the community"])],
     [("Remaining tasks",
       ["*Governance: agreeing microgrid cost sharing among KEPCO, VPP operators, national and local government",
        "*Cooperation with central government: an MOU and task force with the Korea Energy Agency, research on the pilot model, institutionalizing the extended model in national calls"]),
      ("Awaiting decisions",
       ["An integrated coordination system across overlapping departments",
        "The dedicated organization, and the budget and financing method"])]])

slide_closing(prs,
    "From building power plants\nto an economy villages own and the region circulates",
    "Jeonnam-Gwangju Integrated Special City Transition Committee — 'Citizen-Sovereign Community Wealth Building': renewable benefit sharing → Village Salary → community wealth",
    ["Renewable energy is the most stable business model, and therefore the best starting point for community wealth",
     "Residents discover assets (①), enterprises create income (②), domain AI supports them (③), and finance recycles and reinvests (④)",
     "A local government model that extends the national Sunlight Income Village to townships and cities, to microgrids and village electrification, and to a circular economy",
     "To be driven through a public–private governance system centered on the Social and Solidarity Economy Division of the city's Citizen Sovereignty Headquarters"])

prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))
