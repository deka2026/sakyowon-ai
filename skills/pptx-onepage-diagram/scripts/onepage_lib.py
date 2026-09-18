# -*- coding: utf-8 -*-
"""한 장짜리 체계도 PPTX 공통 부품 (python-pptx).

계층(band) × 카드/칩 격자로 1장 다이어그램을 조립한다.
좌표 단위는 인치, 슬라이드는 16:9(13.333 x 7.5).

    from onepage_lib import Deck, INK, BLUE, TEAL, AMBER, PLUM, ROSE, SLATE

    d = Deck()
    d.header("제목", "부제", ["오른쪽 1줄", "오른쪽 2줄"])
    d.band(1.08, 0.86, "1", "지시 주입", "세션마다 자동 적재", BLUE)
    d.chips(1.08, 0.86, [("칩 제목", ["설명1", "설명2"])], BLUE, FILL_BLUE, BORDER_BLUE)
    d.arrow(2.00)
    d.cards(3.12, 1.32, [(AMBER, "카드 제목", ["불릿1", "불릿2"])])
    d.save(r"D:\\...\\파일.pptx")
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

FONT = "맑은 고딕"

INK   = RGBColor(0x14, 0x21, 0x3D)
BLUE  = RGBColor(0x1D, 0x4E, 0xD8)
TEAL  = RGBColor(0x0E, 0x7C, 0x7B)
AMBER = RGBColor(0xB4, 0x53, 0x09)
PLUM  = RGBColor(0x7C, 0x3A, 0xED)
ROSE  = RGBColor(0xBE, 0x18, 0x5D)
SLATE = RGBColor(0x47, 0x55, 0x69)
GREY  = RGBColor(0x64, 0x74, 0x8B)
LINE  = RGBColor(0xCB, 0xD5, 0xE1)
BG    = RGBColor(0xF8, 0xFA, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# 칩 배경/테두리 한 쌍 (색 계열별)
FILL_BLUE,  BORDER_BLUE  = RGBColor(0xEF, 0xF4, 0xFF), RGBColor(0xBF, 0xDB, 0xFE)
FILL_TEAL,  BORDER_TEAL  = RGBColor(0xEC, 0xFD, 0xF5), RGBColor(0x99, 0xF6, 0xE4)
FILL_ROSE,  BORDER_ROSE  = RGBColor(0xFF, 0xF1, 0xF5), RGBColor(0xFB, 0xCF, 0xE8)
FILL_AMBER, BORDER_AMBER = RGBColor(0xFF, 0xF7, 0xED), RGBColor(0xFE, 0xD7, 0xAA)

SLIDE_W, SLIDE_H = 13.333, 7.5


class Deck(object):
    def __init__(self, left=0.55, label_w=1.42, header_h=0.95):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Emu(12192000), Emu(6858000)
        self.s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.LX = left
        self.LW = label_w
        self.CX = left + label_w + 0.14      # 라벨 열을 쓸 때 내용 시작 x
        self.CW = SLIDE_W - self.CX - left   # 라벨 열을 쓸 때 내용 폭
        self.RW = SLIDE_W - left * 2         # 라벨 열이 없을 때 전체 폭
        self.header_h = header_h
        self.rect(0, 0, SLIDE_W, SLIDE_H, fill=BG, shape=MSO_SHAPE.RECTANGLE)

    # ── 기본 도형 ────────────────────────────────────────────────
    def rect(self, x, y, w, h, fill=None, line=None, lw=0.75,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None):
        sh = self.s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if adj is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = adj
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid(); sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line; sh.line.width = Pt(lw)
        sh.shadow.inherit = False          # 기본 그림자를 끄지 않으면 지저분해진다
        sh.text_frame.word_wrap = True
        return sh

    def txt(self, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0):
        """runs = [(텍스트, pt, 색, 굵게), ...] — 한 줄에 한 문단."""
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        self._fill(tf, runs, align, space)
        return tb

    def label(self, shape, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, space=1):
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = Inches(0.08)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        tf.vertical_anchor = anchor
        self._fill(tf, runs, align, space)

    def _fill(self, tf, runs, align, space):
        for i, (t, size, color, bold) in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(space)
            r = p.add_run(); r.text = t
            r.font.name = FONT; r.font.size = Pt(size)
            r.font.color.rgb = color; r.font.bold = bold

    # ── 조립 부품 ────────────────────────────────────────────────
    def header(self, title, subtitle, right_lines=(), h=None):
        h = h or self.header_h
        self.rect(0, 0, SLIDE_W, h, fill=INK, shape=MSO_SHAPE.RECTANGLE)
        self.txt(self.LX, 0.15, 9.2, 0.32, [(title, 24, WHITE, True)])
        self.txt(self.LX, 0.60, 9.6, 0.24,
                 [(subtitle, 11, RGBColor(0xC7, 0xD2, 0xFE), False)])
        if right_lines:
            self.txt(9.60, 0.20, 3.20, 0.60,
                     [(t, 10.5, RGBColor(0xA5, 0xB4, 0xFC), True) for t in right_lines],
                     align=PP_ALIGN.RIGHT, space=2)

    def band(self, y, h, no, name, sub, color):
        """왼쪽 계층 라벨. 같은 y·h로 chips()/cards()를 호출해 짝을 맞춘다."""
        b = self.rect(self.LX, y, self.LW, h, fill=color, adj=0.12)
        self.label(b, [(no + "  " + name, 10.5, WHITE, True),
                       (sub, 8, RGBColor(0xE2, 0xE8, 0xF0), False)],
                   align=PP_ALIGN.CENTER, space=2)

    def chips(self, y, h, data, color, fill, border, tsize=9, dsize=7.5, gap=0.10,
              x=None, w=None):
        """data = [(제목, 설명 or [설명줄, ...]), ...] — 가운데 정렬 작은 상자."""
        x0 = self.CX if x is None else x
        w0 = self.CW if w is None else w
        cw = (w0 - gap * (len(data) - 1)) / len(data)
        for i, (t, d) in enumerate(data):
            c = self.rect(x0 + i * (cw + gap), y, cw, h, fill=fill, line=border)
            ds = d if isinstance(d, (list, tuple)) else [d]
            self.label(c, [(t, tsize, color, True)] + [(u, dsize, SLATE, False) for u in ds],
                       align=PP_ALIGN.CENTER, space=1)

    def cards(self, y, h, data, gap=0.12, tsize=10.5, bsize=8, bullet="▸ ",
              x=None, w=None, top_bar=True):
        """data = [(색, 제목, [불릿, ...]), ...] — 상단 색 띠가 붙은 흰 카드."""
        x0 = self.CX if x is None else x
        w0 = self.CW if w is None else w
        cw = (w0 - gap * (len(data) - 1)) / len(data)
        for i, (color, title, bullets) in enumerate(data):
            cx = x0 + i * (cw + gap)
            self.rect(cx, y, cw, h, fill=WHITE, line=LINE)
            if top_bar:
                self.rect(cx, y, cw, 0.065, fill=color, shape=MSO_SHAPE.RECTANGLE)
            self.txt(cx + 0.13, y + 0.14, cw - 0.26, h - 0.22,
                     [(title, tsize, color, True)]
                     + [(bullet + b, bsize, SLATE, False) for b in bullets], space=2)

    def arrow(self, y, h=0.16, w=0.26, x=None, color=RGBColor(0x94, 0xA3, 0xB8)):
        cx = (self.CX + self.CW / 2 - w / 2) if x is None else x
        a = self.s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(cx), Inches(y),
                                    Inches(w), Inches(h))
        a.fill.solid(); a.fill.fore_color.rgb = color
        a.line.fill.background(); a.shadow.inherit = False

    def strip(self, y, h, text, fill=INK, size=11, adj=0.35, x=None, w=None):
        """맨 아래 한 줄 요약 띠."""
        b = self.rect(x if x is not None else self.LX, y,
                      w if w is not None else self.RW, h, fill=fill, adj=adj)
        self.label(b, [(text, size, WHITE, True)], align=PP_ALIGN.CENTER)

    def save(self, path):
        self.prs.save(path)
        return path
