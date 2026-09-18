# -*- coding: utf-8 -*-
"""예제: AI 업무 체계도 (레이아웃 B — 제목 줄 + 전폭 카드 행 + 하단 요약 띠).

onepage_lib 없이 단독으로 도는 판. 층이 3~4개로 적을 때의 배치를 보여준다.
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = r"D:\사교원 개발그룹\AI업무체계도_20260918.pptx"
F = "맑은 고딕"

INK    = RGBColor(0x14, 0x21, 0x3D)
BLUE   = RGBColor(0x1D, 0x4E, 0xD8)
TEAL   = RGBColor(0x0E, 0x7C, 0x7B)
AMBER  = RGBColor(0xB4, 0x53, 0x09)
PLUM   = RGBColor(0x7C, 0x3A, 0xED)
SLATE  = RGBColor(0x47, 0x55, 0x69)
LINE   = RGBColor(0xCB, 0xD5, 0xE1)
BG     = RGBColor(0xF8, 0xFA, 0xFC)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREY   = RGBColor(0x64, 0x74, 0x8B)

prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
# sldSz type 보정
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
s = prs.slides.add_slide(prs.slide_layouts[6])


def rect(x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if adj is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = adj
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def txt(x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, (t, size, color, bold) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        r = p.add_run(); r.text = t
        r.font.name = F; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    return tb


def label(shape, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.margin_left = Inches(0.10); tf.margin_right = Inches(0.10)
    tf.margin_top = Inches(0.04); tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = anchor
    for i, (t, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(1)
        r = p.add_run(); r.text = t
        r.font.name = F; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold


# ── 배경
bgr = rect(0, 0, 13.333, 7.5, fill=BG, shape=MSO_SHAPE.RECTANGLE)

# ── 헤더
rect(0, 0, 13.333, 0.95, fill=INK, shape=MSO_SHAPE.RECTANGLE)
txt(0.55, 0.16, 9.2, 0.30, [("AI 업무 체계도", 24, WHITE, True)])
txt(0.55, 0.60, 9.6, 0.24,
    [("작업 환경 → 자동화 엔진 → 산출물 → 축적 루프 · 김일영(사회혁신플랫폼) · 2026. 9. 18. 기준", 11, RGBColor(0xC7, 0xD2, 0xFE), False)])
txt(9.60, 0.20, 3.20, 0.60,
    [("스킬 20종 · 기억 30건", 11, RGBColor(0xA5, 0xB4, 0xFC), True),
     ("진행 프로젝트 18건", 11, RGBColor(0xA5, 0xB4, 0xFC), True)],
    align=PP_ALIGN.RIGHT, space=2)

LX, RW = 0.55, 12.23


def band_title(y, no, title, desc, color):
    chip = rect(LX, y, 0.34, 0.30, fill=color, adj=0.25)
    label(chip, [(no, 11, WHITE, True)], align=PP_ALIGN.CENTER)
    txt(LX + 0.46, y + 0.015, 6.0, 0.28, [(title, 13, INK, True)])
    txt(LX + 0.46 + 1.0 + len(title) * 0.125, y + 0.055, 8.0, 0.24, [(desc, 9.5, GREY, False)])


# ══ 1. 기반 환경
band_title(1.12, "1", "작업 환경", "— 한 대의 PC 안에 갖춘 실행 기반", BLUE)
base = rect(LX, 1.50, RW, 0.78, fill=WHITE, line=LINE)
items = [
    ("Claude Code", "데스크톱 · CLI"),
    ("전용 작업폴더", r"D:\사교원 개발그룹 · D:\VideoWorks"),
    ("한글·한셀 COM", "Hwp / HCell 자동화"),
    ("HyperFrames + FFmpeg", "HTML → MP4 렌더"),
    ("GitHub · 자체서버", "deka2026 / Caddy 배포"),
    ("기억(메모리)", "규칙·프로젝트 30건 상시 로드"),
]
cw = (RW - 0.24 - 0.10 * (len(items) - 1)) / len(items)
for i, (t, d) in enumerate(items):
    x = LX + 0.12 + i * (cw + 0.10)
    c = rect(x, 1.61, cw, 0.56, fill=RGBColor(0xEF, 0xF4, 0xFF), line=RGBColor(0xBF, 0xDB, 0xFE))
    label(c, [(t, 10, BLUE, True), (d, 8, SLATE, False)], align=PP_ALIGN.CENTER)


def arrow(x, y, w=0.30, h=0.26, color=RGBColor(0x94, 0xA3, 0xB8)):
    a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False


arrow(6.52, 2.32)

# ══ 2. 자동화 엔진
band_title(2.66, "2", "자동화 엔진 — 반복 업무를 스킬로 고정", "— 한 번 검증한 방법을 스킬 패키지로 만들어 다음부터 명령 한 줄로 실행", TEAL)
engines = [
    (TEAL, "한글문서 자동화",
     ["hwpx-powershell-edit  읽기·수정·표 조립",
      "sujeonghae-routine  “수정해” 교정본 출력",
      "orgchart-doc-migration  조직개편 치환",
      "hwp-binary-extract  구형 .hwp 추출"]),
    (AMBER, "표·예산 자동화",
     ["hancell-budget-report  집행내역 → 통합표",
      "e나라도움 집행완료내역 자동 반영",
      "한셀 COM 재계산·검증 후 .cell 저장",
      "계획 대비 집행률 연도·세목별 산출"]),
    (BLUE, "발표·교육 자료",
     ["pptx-lecture-deck  강의안·발표자료",
      "PowerPoint COM PNG 렌더로 검증",
      "카드형 자동 레이아웃·넘침 교정",
      "컨퍼런스판(15pt·한영 2종) 모드"]),
    (PLUM, "영상·숏폼 제작",
     ["hyperframes 8종  구성·모션·오디오",
      "media-use  음악·이미지·TTS 조달",
      "student-shorts  참가자별 릴스 대량",
      "HTML 코드 → 9:16 MP4 무인 렌더"]),
    (SLATE, "기획·운영·축적",
     ["policy-critique-alternative  비판·대안",
      "mangnam-coop-deploy  수정 → 배포 확인",
      "jeongrihae-routine  “정리해” 4종 자동",
      "핸드오버·스킬·레슨·위키 저장"]),
]
ew = (RW - 0.10 * (len(engines) - 1)) / len(engines)
for i, (c, title, bullets) in enumerate(engines):
    x = LX + i * (ew + 0.10)
    card = rect(x, 3.05, ew, 1.48, fill=WHITE, line=LINE)
    rect(x, 3.05, ew, 0.07, fill=c, shape=MSO_SHAPE.RECTANGLE)
    tb = txt(x + 0.13, 3.19, ew - 0.26, 1.24,
             [(title, 10.5, c, True)] + [("▸ " + b if not b.startswith("  ") else "   " + b.strip(), 8, SLATE, False) for b in bullets],
             space=2)

arrow(6.52, 4.60, h=0.22)

# ══ 3. 산출물
band_title(4.90, "3", "산출물 — 실제로 쓰이고 있는 결과", "— 문서·시스템·영상·지식이 같은 환경에서 나온다", AMBER)
outs = [
    (AMBER, "정책·계획 문서", [
        "시민기금 설립계획(내부 검토용)",
        "고흥 영농형태양광 햇빛소득마을 계획안",
        "에너지공단 공동 시범사업 200MW·92개소",
        "시민주권본부 기능 설계안(4과 14팀·56명)",
        "공론장 발제문 · 서울시의회 행정감사 보고서",
    ]),
    (TEAL, "운영 시스템", [
        "햇소자 — 햇빛소득마을 통합 운영 플랫폼",
        "망남마을협동조합 사이트 · 3교실 신청 백엔드",
        "망남 예산 집행현황 통합문서(2023~2026, 한셀)",
        "연대지능활동가 아카데미 사이트 · 신청 폼",
    ]),
    (PLUM, "교육·미디어", [
        "마을강사단 AI교육 90분 교안 + 강의 PPT 20장",
        "KEA·REN21 국제컨퍼런스 발표자료(19장)",
        "스킴캠프 숏폼 14편(학생 12·강사 2) 무인 렌더",
        "세무사 인터뷰 질문지 등 실무 문서",
    ]),
    (BLUE, "지식 자산", [
        "사교원 위키 — 레슨·업무 매뉴얼 축적",
        "세션 핸드오버 기록(작업·미완료·파일 위치)",
        "재사용 스킬 패키지 20종",
        "문서 규칙 기억(12pt·표 폭·표 우선 등)",
    ]),
]
ow = (RW - 0.12 * (len(outs) - 1)) / len(outs)
for i, (c, title, bullets) in enumerate(outs):
    x = LX + i * (ow + 0.12)
    rect(x, 5.27, ow, 1.32, fill=WHITE, line=LINE)
    rect(x, 5.27, 0.07, 1.32, fill=c, shape=MSO_SHAPE.RECTANGLE)
    txt(x + 0.20, 5.40, ow - 0.34, 1.10,
        [(title, 10.5, c, True)] + [("· " + b, 8, SLATE, False) for b in bullets], space=2)

# ══ 4. 축적 루프
loop = rect(LX, 6.75, RW, 0.48, fill=RGBColor(0x14, 0x21, 0x3D), adj=0.35)
label(loop, [("축적 루프    산출물에서 확인된 규칙을 레슨·스킬·기억으로 되돌려 저장  →  다음 작업은 더 적은 지시로, 더 빠르게  →  다시 산출물",
              11, WHITE, True)], align=PP_ALIGN.CENTER)

prs.save(OUT)
print("saved:", OUT)
